"""Schema-driven tabular observation source with no table-specific aliases.

The physical catalogue is deterministic and cacheable.  A model-backed
Researcher selects real table/column identifiers from that catalogue for each
local task; those selections are execution artifacts, not permanent adapter
code.  CSV, JSON, XLSX, and SQLite use the same navigation contract.
"""

from __future__ import annotations

import csv
import hashlib
from io import BytesIO
import json
import sqlite3
from collections import Counter
from contextlib import closing
from dataclasses import dataclass
from numbers import Real
from pathlib import Path
from statistics import median
from typing import Any, Iterable, Protocol

from .contracts import CoverageProfile, SensorKind, SourceDescriptor
from .data_tool import _aggregate_rows, _handle, _matches_all
from .evidence_contract import observable_manifest_from_rows
from .relational_engine import DuckDBRelationalEngine, RelationTable


CANONICAL_TABLE_FIELDS = (
    "record_id", "paper_id", "sample_id", "source_id",
    "age_ma", "age_error_ma", "latitude", "longitude",
    "position_uncertainty_km", "region", "province", "city", "basin", "locality",
    "rock_type", "formation", "group", "member", "isotope_system", "mineral",
    "technique", "statistical_model", "event_interpretation",
)


class TableBackend(Protocol):
    def table_names(self) -> list[str]: ...
    def columns(self, table_id: str) -> list[str]: ...
    def row_count(self, table_id: str) -> int: ...
    def rows(self, table_id: str) -> Iterable[dict[str, Any]]: ...


@dataclass
class InMemoryTableBackend:
    tables: dict[str, list[dict[str, Any]]]

    def table_names(self) -> list[str]:
        return list(self.tables)

    def columns(self, table_id: str) -> list[str]:
        return list(dict.fromkeys(key for row in self.tables[table_id] for key in row))

    def row_count(self, table_id: str) -> int:
        return len(self.tables[table_id])

    def rows(self, table_id: str) -> Iterable[dict[str, Any]]:
        yield from (dict(row) for row in self.tables[table_id])


class SQLiteTableBackend:
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(f"file:{self.path}?mode=ro", uri=True)
        connection.row_factory = sqlite3.Row
        return connection

    def table_names(self) -> list[str]:
        with closing(self._connect()) as connection:
            return [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]

    def columns(self, table_id: str) -> list[str]:
        self._require_table(table_id)
        with closing(self._connect()) as connection:
            return [row[1] for row in connection.execute(f"PRAGMA table_info({_quote_identifier(table_id)})")]

    def row_count(self, table_id: str) -> int:
        self._require_table(table_id)
        with closing(self._connect()) as connection:
            return int(connection.execute(f"SELECT COUNT(*) FROM {_quote_identifier(table_id)}").fetchone()[0])

    def rows(self, table_id: str) -> Iterable[dict[str, Any]]:
        self._require_table(table_id)
        with closing(self._connect()) as connection:
            for row in connection.execute(f"SELECT * FROM {_quote_identifier(table_id)}"):
                yield dict(row)

    def filtered_rows(
        self,
        table_id: str,
        filters: list[dict[str, Any]],
    ) -> Iterable[dict[str, Any]]:
        """Apply workspace filters inside SQLite before rows cross the data boundary."""

        self._require_table(table_id)
        fields = set(self.columns(table_id))
        clauses: list[str] = []
        parameters: list[Any] = []
        for predicate in filters:
            field = str(predicate.get("field") or "")
            if field not in fields:
                raise ValueError(f"Workspace filter names an unknown field: {field!r}")
            operator = str(predicate.get("operator") or "")
            column = _quote_identifier(field)
            if operator == "is_null":
                clauses.append(f"{column} IS NULL")
            elif operator == "not_null":
                clauses.append(f"{column} IS NOT NULL")
            elif operator == "between":
                clauses.append(f"CAST({column} AS REAL) BETWEEN ? AND ?")
                parameters.extend(
                    [predicate.get("minimum_number"), predicate.get("maximum_number")]
                )
            elif operator == "contains":
                values = [str(item) for item in predicate.get("text_values", [])]
                if not values:
                    clauses.append("0")
                    continue
                clauses.append(
                    "("
                    + " OR ".join(
                        f"LOWER(CAST({column} AS TEXT)) LIKE LOWER(CAST(? AS TEXT))"
                        for _ in values
                    )
                    + ")"
                )
                parameters.extend(f"%{value}%" for value in values)
            elif operator in {"eq", "in"}:
                values = list(predicate.get("text_values", []))
                if not values:
                    clauses.append("0")
                    continue
                clauses.append(
                    "LOWER(CAST({} AS TEXT)) IN ({})".format(
                        column,
                        ", ".join("LOWER(CAST(? AS TEXT))" for _ in values),
                    )
                )
                parameters.extend(values)
            else:
                raise ValueError(f"Unsupported filter operator: {operator!r}")
        where = f" WHERE {' AND '.join(clauses)}" if clauses else ""
        sql = f"SELECT * FROM {_quote_identifier(table_id)}{where}"
        with closing(self._connect()) as connection:
            for row in connection.execute(sql, parameters):
                yield dict(row)

    def _require_table(self, table_id: str) -> None:
        if table_id not in self.table_names():
            raise ValueError(f"Unknown SQLite table: {table_id}")


class DelimitedTableBackend:
    def __init__(self, path: str | Path, *, delimiter: str = ",") -> None:
        self.path, self.delimiter = Path(path), delimiter
        with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter=self.delimiter)
            self._columns = _unique_headers(next(reader, []))
            self._row_count = sum(1 for _ in reader)

    def table_names(self) -> list[str]: return [self.path.stem]
    def columns(self, table_id: str) -> list[str]: self._require(table_id); return list(self._columns)
    def row_count(self, table_id: str) -> int: self._require(table_id); return self._row_count

    def rows(self, table_id: str) -> Iterable[dict[str, Any]]:
        self._require(table_id)
        with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter=self.delimiter)
            next(reader, None)
            for values in reader:
                yield {column: values[index] if index < len(values) else None for index, column in enumerate(self._columns)}

    def _require(self, table_id: str) -> None:
        if table_id != self.path.stem: raise ValueError(f"Unknown delimited table: {table_id}")


class WorkbookTableBackend:
    def __init__(self, path: str | Path, *, header_rows: dict[str, int] | None = None) -> None:
        try:
            import openpyxl  # type: ignore
        except ImportError as error:  # pragma: no cover
            raise RuntimeError("XLSX support requires the optional openpyxl dependency") from error
        self.path, self._openpyxl = Path(path), openpyxl
        workbook = self._load_workbook()
        self._metadata: dict[str, dict[str, Any]] = {}
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            requested = (header_rows or {}).get(sheet_name)
            header_row, headers = self._header(sheet, requested)
            if sheet.max_row is None:
                row_count = sum(
                    1
                    for values in sheet.iter_rows(
                        min_row=header_row + 1, max_row=None, values_only=True
                    )
                    if any(value not in (None, "") for value in values)
                )
            else:
                row_count = max(0, sheet.max_row - header_row)
            self._metadata[sheet_name] = {"header_row": header_row, "columns": headers, "row_count": row_count}
        workbook.close()

    def table_names(self) -> list[str]: return list(self._metadata)
    def columns(self, table_id: str) -> list[str]: return list(self._metadata[table_id]["columns"])
    def row_count(self, table_id: str) -> int: return int(self._metadata[table_id]["row_count"])

    def rows(self, table_id: str) -> Iterable[dict[str, Any]]:
        metadata = self._metadata[table_id]
        workbook = self._load_workbook()
        try:
            sheet = workbook[table_id]
            for values in sheet.iter_rows(min_row=metadata["header_row"] + 1, values_only=True):
                if not any(value not in (None, "") for value in values):
                    continue
                yield {column: values[index] if index < len(values) else None for index, column in enumerate(metadata["columns"])}
        finally:
            workbook.close()

    def _load_workbook(self) -> Any:
        """Load OOXML even when a download has an extra ``.xls`` suffix."""

        if self.path.suffix.casefold() == ".xls":
            with self.path.open("rb") as handle:
                payload = handle.read()
            if payload[:4] == b"PK\x03\x04":
                return self._openpyxl.load_workbook(
                    BytesIO(payload), read_only=True, data_only=True
                )
            raise ValueError(
                "Legacy .xls workbooks are not supported; convert to .xlsx first"
            )
        return self._openpyxl.load_workbook(
            self.path, read_only=True, data_only=True
        )

    @staticmethod
    def _header(sheet: Any, requested: int | None) -> tuple[int, list[str]]:
        if requested is not None:
            values = next(sheet.iter_rows(min_row=requested, max_row=requested, values_only=True), ())
            return requested, _unique_headers(_trim_header_values(values))
        candidates = []
        max_row = min(sheet.max_row, 25) if sheet.max_row is not None else None
        for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=max_row, values_only=True), start=1):
            nonempty = [value for value in values if value not in (None, "")]
            if nonempty:
                candidates.append((len(nonempty), len({str(value) for value in nonempty}), -row_number, row_number, values))
            if row_number >= 25:
                break
        if not candidates:
            return 1, []
        _, _, _, row_number, values = max(candidates)
        return row_number, _unique_headers(_trim_header_values(values))


def _trim_header_values(values: Iterable[Any]) -> list[Any]:
    """Drop formatting-only empty columns after the last physical header."""

    values = list(values)
    last = 0
    for index, value in enumerate(values, start=1):
        if value not in (None, ""):
            last = index
    return values[:last]


class JsonTableBackend(InMemoryTableBackend):
    def __init__(self, path: str | Path) -> None:
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
        if isinstance(payload, list):
            tables = {Path(path).stem: payload}
        elif isinstance(payload, dict) and all(isinstance(value, list) for value in payload.values()):
            tables = payload
        else:
            rows = payload.get("records") or payload.get("rows") or payload.get("data") if isinstance(payload, dict) else None
            if not isinstance(rows, list):
                raise ValueError("JSON table source requires a row list or a map of table names to row lists")
            tables = {Path(path).stem: rows}
        super().__init__({str(name): [dict(row) for row in rows if isinstance(row, dict)] for name, rows in tables.items()})


def backend_from_path(path: str | Path, *, header_rows: dict[str, int] | None = None) -> TableBackend:
    file_path = Path(path)
    suffix = file_path.suffix.casefold()
    if suffix in {".sqlite", ".sqlite3", ".db"}: return SQLiteTableBackend(file_path)
    if suffix == ".csv": return DelimitedTableBackend(file_path)
    if suffix == ".tsv": return DelimitedTableBackend(file_path, delimiter="\t")
    if suffix in {".xlsx", ".xlsm"}: return WorkbookTableBackend(file_path, header_rows=header_rows)
    if suffix == ".xls":
        with file_path.open("rb") as handle:
            if handle.read(4) == b"PK\x03\x04":
                return WorkbookTableBackend(file_path, header_rows=header_rows)
        raise ValueError("Legacy .xls workbooks are not supported; convert to .xlsx first")
    if suffix == ".json": return JsonTableBackend(file_path)
    raise ValueError(f"Unsupported table format: {suffix}")


class DynamicTableObservationSource:
    """A source whose semantic projection is selected per local research task."""

    def __init__(self, backend: TableBackend, *, source_id: str, version: str = "memory", sensor_kind: SensorKind = "event_clock_sensor", name: str | None = None, description: str = "Dynamically navigated read-only tabular observations.", profile_rows: int = 500, max_query_rows: int = 100_000) -> None:
        self.backend, self.source_id, self.version, self.sensor_kind = backend, source_id, version, sensor_kind
        self.name, self.description = name or source_id, description
        self.profile_rows, self.max_query_rows = max(10, profile_rows), max(1, max_query_rows)
        self._catalog = self._build_catalog()

    @classmethod
    def from_path(cls, path: str | Path, *, source_id: str | None = None, sensor_kind: SensorKind = "event_clock_sensor", header_rows: dict[str, int] | None = None, profile_rows: int = 500, max_query_rows: int = 100_000) -> "DynamicTableObservationSource":
        file_path = Path(path)
        stat = file_path.stat()
        version = hashlib.sha256(f"{file_path.resolve()}:{stat.st_size}:{stat.st_mtime_ns}".encode()).hexdigest()[:16]
        return cls(backend_from_path(file_path, header_rows=header_rows), source_id=source_id or file_path.stem, version=version, sensor_kind=sensor_kind, name=file_path.name, profile_rows=profile_rows, max_query_rows=max_query_rows)

    def _build_catalog(self) -> dict[str, Any]:
        tables = []
        for table_id in self.backend.table_names():
            columns = self.backend.columns(table_id)
            samples = []
            row_iterator = iter(self.backend.rows(table_id))
            try:
                for index, row in enumerate(row_iterator):
                    if index >= self.profile_rows: break
                    samples.append(row)
            finally:
                close = getattr(row_iterator, "close", None)
                if close is not None:
                    close()
            profiles = [self._column_profile(column, samples) for column in columns]
            tables.append({"table_id": table_id, "row_count": self.backend.row_count(table_id), "columns": profiles, "sample_rows": samples[:5]})
        return {"catalogue_contract": "tabular_catalogue/v1", "source_id": self.source_id, "version": self.version, "tables": tables}

    @staticmethod
    def _column_profile(column: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
        values = [row.get(column) for row in rows if row.get(column) not in (None, "")]
        numeric = [_number(value) for value in values]
        numeric_values = [value for value in numeric if value is not None]
        inferred = "number" if values and len(numeric_values) == len(values) else ("string" if values else "unknown")
        distinct = list(dict.fromkeys(str(value)[:200] for value in values))[:30]
        return {
            "column_id": column, "inferred_type": inferred,
            "sample_non_null": len(values), "sample_distinct": len({str(value) for value in values}),
            "sample_values": distinct,
            "numeric_range": {"min": min(numeric_values), "max": max(numeric_values)} if numeric_values else None,
        }

    def navigation_catalogue(self) -> dict[str, Any]: return self._catalog

    def describe_source(self) -> SourceDescriptor:
        fields = list(dict.fromkeys(column["column_id"] for table in self._catalog["tables"] for column in table["columns"]))
        return SourceDescriptor(self.source_id, self.name, self.sensor_kind, self.version, self.description, fields, ["dynamic_table_navigation", "read_only_query"], {})

    def describe_scout_summary(self) -> dict[str, Any]:
        return {
            "source_id": self.source_id, "name": self.name, "sensor_kind": self.sensor_kind,
            "version": self.version, "description": self.description,
            "variables": [f"{len(self._catalog['tables'])} physical table(s); columns are inspected only during Research"],
            "capabilities": ["dynamic_table_navigation", "read_only_query"], "authoritative": True,
        }

    def describe_schema(self) -> dict[str, Any]: return self._catalog
    def describe_capabilities(self) -> list[str]: return ["dynamic_table_navigation", "read_only_query"]
    def audit_semantic_bindings(self) -> dict[str, Any]: return {"status": "bound_per_research_action", "schema_fingerprint": self.version, "bindings": []}
    def describe_query_contract(self) -> dict[str, Any]: return {"query_contract": "dynamic_table_query/v1", **self._catalog}

    def workspace_catalogue(self) -> dict[str, Any]:
        """Expose physical table fields for local, read-only data operations.

        The scientific semantic projection remains a separate Researcher
        decision.  This catalogue merely lets the same Researcher count,
        sample, filter, and group a table before committing to that projection.
        """

        operations = ["inspect", "profile", "sample", "query", "aggregate", "materialize"]
        if len(self._catalog["tables"]) >= 2:
            operations.append("relate")
        return {
            "source_id": self.source_id,
            "data_operation_contract": "data_operation/v1",
            "decision_contract": "data_decision_required/v1",
            "tables": [
                {"table_id": table["table_id"], "fields": [column["column_id"] for column in table["columns"]]}
                for table in self._catalog["tables"]
            ],
            "operations": operations,
            "aggregate_operations": ["count", "min", "max", "mean", "median", "std", "n_unique"],
            "relation_contract": {
                "status": "available" if len(self._catalog["tables"]) >= 2 else "not_applicable",
                "key_selection": "explicit physical columns selected from the live table catalogue",
                "modes": ["probe", "materialize"] if len(self._catalog["tables"]) >= 2 else [],
                "matching_policy": "must_be_explicit_in_researcher_request",
            },
        }

    def workspace_execute(self, request: dict[str, Any]) -> dict[str, Any]:
        operation, table_id = str(request.get("operation") or ""), str(request.get("table_id") or "")
        # Relation shape is authoritative.  Accepting materialize here keeps
        # old checkpoints resumable, while canonicalizing the request into
        # the same relation implementation instead of silently treating it
        # as a single-table query.
        has_graph = bool(request.get("tables") or request.get("relations"))
        has_pair = bool(request.get("related_table_id"))
        if operation == "relate" or (operation == "materialize" and (has_graph or has_pair)):
            relation_request = dict(request)
            relation_request["operation"] = "relate"
            relation_request.setdefault("relation_mode", "materialize" if operation == "materialize" else "probe")
            decision = self._relation_policy_decision(relation_request)
            if decision is not None:
                return decision
            if request.get("tables") or request.get("relations"):
                return self._relate_table_graph(relation_request)
            return self._relate_tables(relation_request)
        table = next((item for item in self._catalog["tables"] if item["table_id"] == table_id), None)
        if table is None:
            raise ValueError(f"Unknown table for data tool: {table_id!r}")
        fields = {column["column_id"] for column in table["columns"]}
        for field in [*request.get("columns", []), *request.get("group_by", []), str(request.get("bin_field") or "")]:
            if field and field not in fields:
                raise ValueError(f"Workspace request names an unknown field: {field!r}")
        for predicate in request.get("filters", []):
            if predicate.get("field") not in fields:
                raise ValueError(f"Workspace filter names an unknown field: {predicate.get('field')!r}")
        for aggregate in request.get("aggregates", []):
            field = aggregate.get("field")
            if field and field not in fields:
                raise ValueError(f"Workspace aggregate names an unknown field: {field!r}")
        if operation == "inspect":
            return {"summary": f"Inspected table {table_id} from {self.source_id}.", "findings": {"catalogue": self.workspace_catalogue(), "table": table}}
        columns = list(fields)
        rows, execution_request = self._workspace_rows(table_id, request)
        engine = DuckDBRelationalEngine()
        if operation == "profile":
            result = engine.execute_table(rows, columns, execution_request)
            return {"summary": f"Profiled {result['record_count']:,} matching rows from table {table_id} with DuckDB.", "findings": result}
        if operation in {"sample", "query"}:
            result = engine.execute_table(rows, columns, execution_request)
            return {"summary": f"Read {result['returned_records']:,} of {result['record_count']:,} matching rows from table {table_id} with DuckDB.", "findings": result}
        if operation == "materialize":
            result = engine.execute_table(rows, columns, execution_request)
            handle = _handle(self.source_id, request, result["record_count"])
            handle["table_id"] = table_id
            return {"summary": f"Created a durable query handle for {result['record_count']:,} rows from table {table_id}; DuckDB executed the selection and rows were not placed in model context.", "findings": {"record_count": result["record_count"], "dataset_handle": handle}}
        if operation == "aggregate":
            result = engine.execute_table(rows, columns, execution_request)
            return {"summary": f"Aggregated {result['input_record_count']:,} records from {self.source_id} into {result['group_count']:,} group(s) with DuckDB.", "findings": result}
        raise ValueError(f"Unsupported dynamic-table workspace operation: {operation!r}")

    def _workspace_rows(
        self,
        table_id: str,
        request: dict[str, Any],
    ) -> tuple[Iterable[dict[str, Any]], dict[str, Any]]:
        """Push filters into a capable backend and remove them from the later engine pass."""

        filters = [
            dict(item)
            for item in request.get("filters", [])
            if isinstance(item, dict)
        ]
        filtered_rows = getattr(self.backend, "filtered_rows", None)
        if filters and callable(filtered_rows):
            execution_request = {**request, "filters": []}
            return filtered_rows(table_id, filters), execution_request
        return self.backend.rows(table_id), request

    def _prepare_relation_graph(self, request: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
        table_specs = request.get("tables") or []
        raw_edges = request.get("relations") or []
        if not isinstance(table_specs, list) or not isinstance(raw_edges, list) or len(table_specs) < 2 or not raw_edges:
            raise ValueError("A multi-table relation requires at least two tables and one relation edge")
        known_tables = {table["table_id"] for table in self._catalog["tables"]}
        normalized: dict[str, dict[str, Any]] = {}
        aliases: set[str] = set()
        for item in table_specs:
            if not isinstance(item, dict):
                raise ValueError("Relation graph table entries must be objects")
            table_id = str(item.get("table_id") or "")
            alias = str(item.get("alias") or table_id)
            if table_id not in known_tables or table_id in normalized:
                raise ValueError(f"Relation graph names an unknown or duplicate table: {table_id!r}")
            if not alias or alias in aliases:
                raise ValueError("Relation graph table aliases must be unique and non-empty")
            normalized[table_id] = {"table_id": table_id, "alias": alias, "columns": self.backend.columns(table_id)}
            aliases.add(alias)

        # A relation may filter only its declared primary table.  The old
        # implementation stored ``filters`` in the durable request but built
        # the relation from both complete physical tables, so a bounded query
        # silently became an all-record join.  Make the ownership explicit
        # rather than guessing which table a physical field belongs to.
        filters = request.get("filters") or []
        filter_table_id = str(request.get("table_id") or "")
        if filters:
            if filter_table_id not in normalized:
                raise ValueError("Relation filters require table_id to name one declared relation table")
            filter_fields = set(normalized[filter_table_id]["columns"])
            for predicate in filters:
                if not isinstance(predicate, dict) or str(predicate.get("field") or "") not in filter_fields:
                    raise ValueError("Relation filters must name physical columns of the declared table_id")

        edges: list[dict[str, Any]] = []
        for item in raw_edges:
            if not isinstance(item, dict):
                raise ValueError("Relation graph edges must be objects")
            left_table = str(item.get("left_table_id") or "")
            right_table = str(item.get("right_table_id") or "")
            left_key = str(item.get("left_key") or "")
            right_key = str(item.get("right_key") or "")
            join_type = str(item.get("join_type") or "inner")
            match_policy = str(item.get("match_policy") or request.get("match_policy") or "exact")
            match_tolerance = float(item.get("match_tolerance") or request.get("match_tolerance") or 0)
            if left_table not in normalized or right_table not in normalized or left_table == right_table:
                raise ValueError("Relation graph edges must connect two distinct declared tables")
            if left_key not in set(normalized[left_table]["columns"]) or right_key not in set(normalized[right_table]["columns"]):
                raise ValueError("Relation graph edges must use physical columns from their declared tables")
            if join_type not in {"inner", "left"}:
                raise ValueError("Multi-table relation graphs support only inner and left joins")
            edges.append({"left_table_id": left_table, "left_key": left_key, "right_table_id": right_table, "right_key": right_key, "join_type": join_type, "match_policy": match_policy, "match_tolerance": match_tolerance})
        if not edges:
            raise ValueError("Relation graph requires at least one edge")

        root = next(iter(normalized))
        visited = {root}
        steps: list[dict[str, Any]] = []
        pending = list(edges)
        while pending:
            progress = False
            remaining: list[dict[str, Any]] = []
            for edge in pending:
                left, right = edge["left_table_id"], edge["right_table_id"]
                if left in visited and right not in visited:
                    steps.append({**edge, "existing_table_id": left, "existing_key": edge["left_key"], "new_table_id": right, "new_key": edge["right_key"]})
                    visited.add(right); progress = True
                elif right in visited and left not in visited:
                    steps.append({**edge, "existing_table_id": right, "existing_key": edge["right_key"], "new_table_id": left, "new_key": edge["left_key"]})
                    visited.add(left); progress = True
                elif left in visited and right in visited:
                    raise ValueError("Relation graph contains a cycle; use a tree of explicit join edges")
                else:
                    remaining.append(edge)
            if not progress:
                raise ValueError("Relation graph is disconnected from its first table")
            pending = remaining
        return normalized, edges, steps

    def _relation_rows(self, request: dict[str, Any], table_id: str) -> tuple[dict[str, Any], ...]:
        """Return the physical rows for one relation input under its declared scope."""

        filters = request.get("filters") or []
        if not filters:
            return tuple(dict(row) for row in self.backend.rows(table_id))
        filter_table_id = str(request.get("table_id") or "")
        if filter_table_id not in set(self.backend.table_names()):
            raise ValueError("Relation filters require table_id to name a live relation table")
        if table_id != filter_table_id:
            return tuple(dict(row) for row in self.backend.rows(table_id))
        fields = set(self.backend.columns(table_id))
        if any(not isinstance(predicate, dict) or str(predicate.get("field") or "") not in fields for predicate in filters):
            raise ValueError("Relation filters must name physical columns of the declared table_id")
        rows, execution_request = self._workspace_rows(table_id, request)
        if not execution_request.get("filters"):
            return tuple(dict(row) for row in rows)
        return tuple(row for row in rows if _matches_all(row, filters))

    def _relate_table_graph(self, request: dict[str, Any]) -> dict[str, Any]:
        tables, edges, steps = self._prepare_relation_graph(request)
        relation_tables = [
            RelationTable(
                table_id=table_id,
                alias=str(spec["alias"]),
                columns=tuple(spec["columns"]),
                rows=self._relation_rows(request, table_id),
            )
            for table_id, spec in tables.items()
        ]
        relation_mode = str(request.get("relation_mode") or "probe")
        if relation_mode not in {"probe", "materialize"}:
            raise ValueError(f"Unsupported relation mode: {relation_mode!r}")
        row_limit = max(1, min(int(request.get("row_limit") or self.max_query_rows), self.max_query_rows))
        engine = DuckDBRelationalEngine()
        edge_reports, joined_count, overflow = engine.relation_probe_or_count(relation_tables, steps, edges, row_limit)
        findings = {
            "relation_contract": "table_relation_graph/v1",
            "relation_mode": relation_mode,
            "tables": [{"table_id": item["table_id"], "alias": item["alias"]} for item in tables.values()],
            "relations": edges,
            "join_order": steps,
            "edge_reports": edge_reports,
            "connected": True,
            "acyclic": True,
            "estimated_joined_row_count": None if overflow else joined_count,
            "estimated_exceeds_declared_limit": overflow,
        }
        if request.get("filters"):
            findings["relation_scope"] = {
                "filtered_table_id": str(request.get("table_id") or ""),
                "filters": list(request["filters"]),
                "filtered_input_record_count": next(
                    len(table.rows) for table in relation_tables
                    if table.table_id == str(request.get("table_id") or "")
                ),
            }
        if relation_mode == "probe":
            return {
                "summary": f"Probed a connected {len(tables)}-table relation graph with {len(edges)} join edge(s).",
                "findings": findings,
                "limitations": ["The probe reports deterministic key coverage and join expansion; it does not assert scientific equivalence of the linked fields.", "No joined rows were materialized."],
            }
        if overflow:
            return {
                "status": "insufficient",
                "summary": f"The multi-table relation exceeds the declared row limit of {row_limit:,}; no handle was created.",
                "findings": findings,
                "limitations": ["Reduce the relation scope or raise the declared row limit after reviewing the per-edge reports."],
            }
        handle = _handle(self.source_id, request, joined_count)
        handle.update({
            "handle_contract": "data_handle/v1",
            "relation_contract": "table_relation_graph/v1",
            "tables": findings["tables"],
            "relations": edges,
            "join_order": steps,
            "relation_mode": "materialize",
        })
        findings["dataset_handle"] = handle
        return {
            "summary": f"Created a durable handle for the {len(tables)}-table relation graph with {joined_count:,} joined rows; DuckDB executed the relation.",
            "findings": findings,
            "limitations": ["The joined rows remain local; downstream tools must consume the returned handle."],
        }

    def _relate_tables(self, request: dict[str, Any]) -> dict[str, Any]:
        """Probe or materialize an explicit relation between two live tables.

        The Researcher chooses the physical tables and key columns.  This
        adapter only checks the relation and records deterministic coverage;
        it never guesses that similarly named columns are equivalent.
        """

        left_id = str(request.get("table_id") or "")
        right_id = str(request.get("related_table_id") or "")
        left_key = str(request.get("left_key") or "")
        right_key = str(request.get("right_key") or "")
        join_type = str(request.get("join_type") or "inner")
        match_policy = str(request.get("match_policy") or "exact")
        match_tolerance = float(request.get("match_tolerance") or 0)
        relation_mode = str(request.get("relation_mode") or "probe")
        table_ids = {table["table_id"] for table in self._catalog["tables"]}
        if left_id not in table_ids or right_id not in table_ids or left_id == right_id:
            raise ValueError("A relation requires two distinct tables from the live catalogue")
        if join_type not in {"inner", "left", "semi", "anti"}:
            raise ValueError(f"Unsupported relation join type: {join_type!r}")
        if relation_mode not in {"probe", "materialize"}:
            raise ValueError(f"Unsupported relation mode: {relation_mode!r}")
        left_fields = set(self.backend.columns(left_id))
        right_fields = set(self.backend.columns(right_id))
        if left_key not in left_fields or right_key not in right_fields:
            raise ValueError("Relation keys must be physical columns of their declared tables")

        relation_tables = [
            RelationTable(left_id, "left", tuple(self.backend.columns(left_id)), self._relation_rows(request, left_id)),
            RelationTable(right_id, "right", tuple(self.backend.columns(right_id)), self._relation_rows(request, right_id)),
        ]
        edge = {"left_table_id": left_id, "left_key": left_key, "right_table_id": right_id, "right_key": right_key, "join_type": join_type, "match_policy": match_policy, "match_tolerance": match_tolerance}
        steps = [{"existing_table_id": left_id, "existing_key": left_key, "new_table_id": right_id, "new_key": right_key, "join_type": join_type, "match_policy": match_policy, "match_tolerance": match_tolerance}]
        engine = DuckDBRelationalEngine()
        fuzzy_report = None
        if match_policy == "fuzzy":
            threshold = float(match_tolerance or 0.85)
            if not 0.0 < threshold <= 1.0:
                raise ValueError("Fuzzy match_tolerance must be a score threshold in (0, 1]")
            fuzzy_report = engine.fuzzy_match_pair(relation_tables[0], relation_tables[1], left_key, right_key, threshold=threshold, row_limit=self.max_query_rows)
            pair_count = int(fuzzy_report["accepted_pair_count"])
            overflow = pair_count > self.max_query_rows
            stats = {**edge, "matched_distinct_keys": pair_count, "joined_pair_count": pair_count, "accepted_pair_count": pair_count, "candidate_pair_count": fuzzy_report["candidate_pair_count"], "ambiguous_pair_count": fuzzy_report["ambiguous_pair_count"], "left_row_count": len(relation_tables[0].rows), "right_row_count": len(relation_tables[1].rows)}
        else:
            edge_reports, pair_count, overflow = engine.relation_probe_or_count(relation_tables, steps, [edge], self.max_query_rows)
            stats = edge_reports[0]
        findings = {
            "relation_contract": "table_relation/v1",
            "left_table_id": left_id,
            "right_table_id": right_id,
            "left_key": left_key,
            "right_key": right_key,
            "join_type": join_type,
            "match_policy": match_policy,
            "match_tolerance": match_tolerance,
            "relation_mode": relation_mode,
            **stats,
        }
        if request.get("filters"):
            findings["relation_scope"] = {
                "filtered_table_id": left_id,
                "filters": list(request["filters"]),
                "filtered_input_record_count": len(relation_tables[0].rows),
            }
        if fuzzy_report is not None:
            findings.update({"match_provider": fuzzy_report["provider"], "accepted_pairs": fuzzy_report["accepted_pairs"], "ambiguous_pairs": fuzzy_report["ambiguous_pairs"]})
        if relation_mode == "probe":
            return {
                "summary": f"Probed {left_id}.{left_key} against {right_id}.{right_key}: {stats['matched_distinct_keys']:,} distinct key matches.",
                "findings": findings,
                "limitations": [
                    "A relation probe validates key coverage and duplicate structure; it does not assert that the fields have the same scientific meaning.",
                    "No joined rows were materialized.",
                ],
            }

        row_limit = max(1, min(int(request.get("row_limit") or self.max_query_rows), self.max_query_rows))
        if pair_count > row_limit:
            return {
                "status": "insufficient",
                "summary": f"The relation would produce {pair_count:,} joined rows, exceeding the declared limit of {row_limit:,}; no handle was created.",
                "findings": findings,
                "limitations": ["Reduce the relation scope or raise the declared row limit after reviewing the probe result."],
            }
        handle = _handle(self.source_id, request, pair_count)
        handle.update({
            "handle_contract": "data_handle/v1",
            "relation_contract": "table_relation/v1",
            "table_id": left_id,
            "related_table_id": right_id,
            "left_key": left_key,
            "right_key": right_key,
            "join_type": join_type,
            "match_policy": match_policy,
            "match_tolerance": match_tolerance,
            "relation_mode": "materialize",
        })
        if fuzzy_report is not None:
            handle["match_provider"] = fuzzy_report["provider"]
            handle["accepted_pairs"] = fuzzy_report["accepted_pairs"]
        findings["dataset_handle"] = handle
        return {
            "summary": f"Created a durable relation handle for {pair_count:,} joined rows from {left_id} and {right_id}.",
            "findings": findings,
            "limitations": ["The joined rows remain local; downstream tools must consume the returned handle."],
        }

    def load_workspace_handle(self, handle: dict[str, Any]) -> list[dict[str, Any]]:
        if handle.get("source_id") != self.source_id:
            raise ValueError("Workspace handle belongs to a different source")
        # A materialized table may have been selected through the dynamic
        # navigation contract.  Reuse that exact binding to return canonical
        # scientific rows; never make downstream tools interpret physical
        # column names a second time.
        projection_scope = handle.get("scope")
        if (
            isinstance(projection_scope, dict)
            and projection_scope.get("_scope_contract") == "table_query/v1"
            and projection_scope.get("bindings")
        ):
            table_id = str(handle.get("table_id") or (handle.get("request") or {}).get("table_id") or "")
            if str(projection_scope.get("table_id") or "") != table_id:
                raise ValueError("Materialized handle projection names a different physical table")
            return self.query(projection_scope)
        if handle.get("relation_contract") == "table_relation/v1":
            return self._load_relation_handle(handle)
        if handle.get("relation_contract") == "table_relation_graph/v1":
            return self._load_relation_graph_handle(handle)
        request = dict(handle.get("request") or {})
        table_id = str(handle.get("table_id") or request.get("table_id") or "")
        table = next((item for item in self._catalog["tables"] if item["table_id"] == table_id), None)
        if table is None:
            raise ValueError("Workspace handle names an unknown table")
        fields = [str(item) for item in request.get("columns", []) if item]
        known = {column["column_id"] for column in table["columns"]}
        if not fields or not set(fields).issubset(known):
            raise ValueError("Workspace handle has invalid dynamic-table columns")
        load_request = {**request, "operation": "query", "columns": fields, "row_limit": max(1, min(int(handle.get("record_count") or self.max_query_rows), self.max_query_rows))}
        rows, execution_request = self._workspace_rows(table_id, load_request)
        result = DuckDBRelationalEngine().execute_table(rows, list(known), execution_request)
        if result["truncated"]:
            raise ValueError("Workspace handle expanded beyond its declared row limit")
        return result["records"]

    def _load_relation_handle(self, handle: dict[str, Any]) -> list[dict[str, Any]]:
        left_id = str(handle.get("table_id") or "")
        right_id = str(handle.get("related_table_id") or "")
        left_key = str(handle.get("left_key") or "")
        right_key = str(handle.get("right_key") or "")
        join_type = str(handle.get("join_type") or "inner")
        if left_id == right_id or left_id not in self.backend.table_names() or right_id not in self.backend.table_names():
            raise ValueError("Relation handle names unknown or identical tables")
        if left_key not in set(self.backend.columns(left_id)) or right_key not in set(self.backend.columns(right_id)):
            raise ValueError("Relation handle names unknown key columns")
        request = dict(handle.get("request") or {})
        row_limit = max(1, min(int(request.get("row_limit") or self.max_query_rows), self.max_query_rows))
        engine = DuckDBRelationalEngine()
        if str(handle.get("match_policy") or "exact") == "fuzzy":
            joined = engine.load_fuzzy_pairs(
                RelationTable(left_id, "left", tuple(self.backend.columns(left_id)), self._relation_rows(request, left_id)),
                RelationTable(right_id, "right", tuple(self.backend.columns(right_id)), self._relation_rows(request, right_id)),
                list(handle.get("accepted_pairs") or []), row_limit,
            )
            if len(joined) > row_limit:
                raise ValueError("Fuzzy relation handle expanded beyond its declared row limit")
            return joined
        joined = engine.load_pair(
            RelationTable(left_id, "left", tuple(self.backend.columns(left_id)), self._relation_rows(request, left_id)),
            RelationTable(right_id, "right", tuple(self.backend.columns(right_id)), self._relation_rows(request, right_id)),
            left_key, right_key, join_type, row_limit, str(handle.get("match_policy") or "exact"),
        )
        if len(joined) > row_limit:
            raise ValueError("Relation handle expanded beyond its declared row limit")
        return joined

    @staticmethod
    def _relation_policy_decision(request: dict[str, Any]) -> dict[str, Any] | None:
        """Refuse implicit semantic changes before any rows are joined."""

        policies = []
        for item in request.get("relations") or []:
            if isinstance(item, dict):
                policies.append(str(item.get("match_policy") or request.get("match_policy") or "exact"))
        if request.get("related_table_id"):
            policies.append(str(request.get("match_policy") or "exact"))
        policies = list(dict.fromkeys(policies or ["exact"]))
        unsupported = [item for item in policies if item not in {"exact", "normalized"}]
        if not unsupported:
            return None
        authorization = str(request.get("semantic_authorization") or "")
        if authorization != "explicit":
            return {
                "operation_status": "decision_required",
                "summary": "The requested relation would change scientific meaning and has no explicit Researcher authorization.",
                "findings": {"decision": {"contract": "data_decision_required/v1", "decision_kind": "matching_policy", "requested_policies": unsupported, "candidates": ["authorize the stated policy", "choose an exact or normalized rule", "reject the relation"], "evidence": ["physical table keys", "candidate relation coverage"]}},
                "limitations": ["No rows were joined.", "The provider does not infer fuzzy, spatial, or temporal matching policy."],
            }
        if request.get("tables") or request.get("relations"):
            return {
                "operation_status": "unavailable",
                "status": "data_unavailable",
                "summary": "The registered fuzzy provider currently supports one explicit two-table relation at a time; the global relation contract remains unchanged.",
                "findings": {"decision": {"contract": "data_decision_required/v1", "decision_kind": "provider_capability", "requested_policies": unsupported, "candidates": ["execute fuzzy relations one edge at a time", "choose exact or normalized policies for the graph"], "evidence": []}},
                "limitations": ["No rows were joined.", "The Researcher must not replace a multi-edge fuzzy relation with an exact relation."],
            }
        non_string_policies = [item for item in unsupported if item in {"spatial", "temporal"}]
        if non_string_policies:
            return {
                "operation_status": "unavailable",
                "status": "data_unavailable",
                "summary": f"The requested {non_string_policies[0]} relation belongs to a spatial/temporal matching provider, not the string entity resolver.",
                "findings": {"decision": {"contract": "data_decision_required/v1", "decision_kind": "provider_capability", "requested_policies": non_string_policies, "candidates": ["register the matching provider", "express the relation with an explicit spatial/temporal tool"], "evidence": []}},
                "limitations": ["No rows were joined.", "The provider does not substitute string similarity for spatial or temporal tolerance."],
            }
        # A two-table request is handled by the registered fuzzy provider
        # inside the same data.tool operation.
        return None

    def _load_relation_graph_handle(self, handle: dict[str, Any]) -> list[dict[str, Any]]:
        request = dict(handle.get("request") or {})
        request.setdefault("tables", handle.get("tables") or [])
        request.setdefault("relations", handle.get("relations") or [])
        tables, _, steps = self._prepare_relation_graph(request)
        row_limit = max(1, min(int((handle.get("request") or {}).get("row_limit") or self.max_query_rows), self.max_query_rows))
        relation_tables = [RelationTable(table_id, str(spec["alias"]), tuple(spec["columns"]), self._relation_rows(request, table_id)) for table_id, spec in tables.items()]
        joined = DuckDBRelationalEngine().load_relation(relation_tables, steps, row_limit=row_limit)
        if len(joined) > row_limit:
            raise ValueError("Multi-table relation handle expanded beyond its declared row limit")
        return joined[:row_limit]

    def describe_discovery_coverage(self) -> dict[str, Any]:
        return {
            "record_count": sum(table["row_count"] for table in self._catalog["tables"]),
            "table_count": len(self._catalog["tables"]),
            "tables": [{"table_id": table["table_id"], "row_count": table["row_count"], "column_count": len(table["columns"])} for table in self._catalog["tables"]],
            "semantic_status": "unbound_until_research",
        }

    def canonicalize_navigation(self, proposal: dict[str, Any]) -> dict[str, Any]:
        table_id = str(proposal.get("table_id") or "")
        table = next((item for item in self._catalog["tables"] if item["table_id"] == table_id), None)
        if table is None: raise ValueError(f"Navigation selected an unknown table: {table_id}")
        columns = {item["column_id"]: item for item in table["columns"]}
        bindings, seen_roles = [], set()
        for item in proposal.get("bindings", []):
            role, column = item.get("canonical_field"), item.get("source_column")
            if role not in CANONICAL_TABLE_FIELDS or column not in columns or role in seen_roles:
                raise ValueError("Navigation contains an unknown or duplicate semantic binding")
            seen_roles.add(role); bindings.append({"canonical_field": role, "source_column": column})
        predicates = []
        for item in proposal.get("predicates", []):
            column, operator = item.get("source_column"), item.get("operator")
            if column not in columns or operator not in {
                "eq", "in", "between", "contains", "is_null", "not_null"
            }:
                raise ValueError("Navigation contains an unsupported predicate")
            value = item.get("values") if "values" in item else item.get("value")
            if operator == "between": value = [item.get("minimum_value"), item.get("maximum_value")]
            predicates.append({"source_column": column, "operator": operator, "value": value})
        status = proposal.get("status")
        if status not in {"resolved", "partial", "scope_unresolvable", "unconstrained"}: raise ValueError("Unknown navigation status")
        if status in {"resolved", "unconstrained"} and not bindings: raise ValueError("Executable table navigation requires at least one semantic binding")
        return {
            "_scope_contract": "table_query/v1", "source_id": self.source_id, "table_id": table_id,
            "status": status, "bindings": bindings, "predicates": predicates,
            "unresolved_terms": [str(item) for item in proposal.get("unresolved_terms", [])],
            "non_filter_terms": [str(item) for item in proposal.get("research_terms_outside_query", [])],
            "limit": min(int(proposal.get("limit") or self.max_query_rows), self.max_query_rows),
            "provenance": {"owner": "researcher_data_navigation", "rationale": str(proposal.get("rationale", ""))[:1000]},
        }

    def identity_navigation(self, requested_scope: dict[str, Any] | None = None) -> dict[str, Any]:
        """Deterministic fallback only when physical names already are canonical."""
        table = self._catalog["tables"][0]
        columns = {item["column_id"] for item in table["columns"]}
        bindings = [{"canonical_field": field, "source_column": field} for field in CANONICAL_TABLE_FIELDS if field in columns]
        return self.canonicalize_navigation({"table_id": table["table_id"], "status": "unconstrained" if bindings else "scope_unresolvable", "bindings": bindings, "predicates": [], "unresolved_terms": [] if bindings else ["No exact canonical column names are present and no navigation advisor is configured."], "research_terms_outside_query": [], "rationale": "Exact-name fallback without alias inference.", "limit": self.max_query_rows})

    def query(self, scope: dict[str, Any]) -> list[dict[str, Any]]:
        if scope.get("_scope_contract") != "table_query/v1" or scope.get("status") not in {"resolved", "partial", "unconstrained"}:
            return []
        selected = []
        for row in self.backend.rows(scope["table_id"]):
            if not all(_matches(row, predicate) for predicate in scope.get("predicates", [])): continue
            selected.append(self._project(row, scope["bindings"], len(selected) + 1))
            if len(selected) > scope.get("limit", self.max_query_rows):
                raise ValueError("Dynamic table query exceeded its declared row limit")
        return selected

    def _project(self, row: dict[str, Any], bindings: list[dict[str, str]], index: int) -> dict[str, Any]:
        projected = {item["canonical_field"]: row.get(item["source_column"]) for item in bindings}
        for field in ("age_ma", "age_error_ma", "latitude", "longitude", "position_uncertainty_km"):
            if field in projected: projected[field] = _number(projected[field])
        projected.setdefault("record_id", f"{self.source_id}:{index}")
        projected.setdefault("source_id", self.source_id)
        projected["_semantic_field_provenance"] = {item["canonical_field"]: [item["source_column"]] for item in bindings}
        projected["position"] = {
            "latitude": projected.get("latitude"), "longitude": projected.get("longitude"),
            "uncertainty_km": projected.get("position_uncertainty_km"),
            "resolution": "unresolved", "provenance": "dynamically bound table columns",
        }
        return projected

    def sample_records(self, scope: dict[str, Any], n: int = 5) -> list[dict[str, Any]]: return self.query({**scope, "limit": min(n, scope.get("limit", n))})[:n]
    def estimate_query(self, scope: dict[str, Any]) -> dict[str, Any]: return {"estimated_records": None, "estimated_cost": "requires_research_query", "query_scope": scope}

    def describe_observables(self, scope: dict[str, Any]) -> dict[str, dict[str, Any]]:
        return observable_manifest_from_rows(self.query(scope)) if scope.get("_scope_contract") == "table_query/v1" else {}

    def describe_resolution(self, scope: dict[str, Any]) -> dict[str, Any]:
        rows = self.query(scope)
        positioned = [row for row in rows if row.get("latitude") is not None and row.get("longitude") is not None]
        errors = [row.get("position_uncertainty_km") for row in positioned if row.get("position_uncertainty_km") is not None]
        return {"coordinate_coverage": len(positioned) / len(rows) if rows else 0.0, "median_spatial_uncertainty_km": median(errors) if errors else None, "minimum_resolvable_migration_km": median(errors) * 2 if errors else None}

    def describe_uncertainty(self, scope: dict[str, Any]) -> dict[str, Any]:
        rows = self.query(scope); errors = [row["age_error_ma"] for row in rows if row.get("age_error_ma") is not None]
        return {"age_error_coverage": len(errors) / len(rows) if rows else 0.0, "median_age_error_ma": median(errors) if errors else None, "median_spatial_uncertainty_km": self.describe_resolution(scope).get("median_spatial_uncertainty_km")}

    def profile_coverage(self, scope: dict[str, Any]) -> CoverageProfile:
        rows = self.query(scope)
        tracked = ("age_ma", "age_error_ma", "latitude", "longitude", "rock_type", "source_id", "event_interpretation", "mineral", "technique")
        missing = {field: sum(row.get(field) in (None, "") for row in rows) / len(rows) if rows else 1.0 for field in tracked}
        return CoverageProfile(self.source_id, len(rows), sum(table["row_count"] for table in self._catalog["tables"]), {"observables": observable_manifest_from_rows(rows)}, missing, self.describe_resolution(scope), self.describe_uncertainty(scope), self.estimate_query(scope), rows[:5], [], self.version)


def _number(value: Any) -> float | None:
    if value in (None, ""): return None
    if isinstance(value, Real): return float(value)
    try: return float(str(value).strip())
    except (TypeError, ValueError): return None


def _relation_key(value: Any) -> str | None:
    """Return a stable join key while treating null/blank values as absent."""

    if value in (None, ""):
        return None
    return str(value).strip()


def _relation_statistics(
    left_rows: list[dict[str, Any]],
    right_rows: list[dict[str, Any]],
    left_key: str,
    right_key: str,
) -> tuple[dict[str, Any], dict[str, list[dict[str, Any]]]]:
    right_index: dict[str, list[dict[str, Any]]] = {}
    for row in right_rows:
        key = _relation_key(row.get(right_key))
        if key is not None:
            right_index.setdefault(key, []).append(row)
    left_index: dict[str, int] = {}
    for row in left_rows:
        key = _relation_key(row.get(left_key))
        if key is not None:
            left_index[key] = left_index.get(key, 0) + 1
    matched_keys = sorted(set(left_index).intersection(right_index))
    matched_left_rows = sum(left_index[key] for key in matched_keys)
    matched_right_rows = sum(len(right_index[key]) for key in matched_keys)
    joined_pair_count = sum(left_index[key] * len(right_index[key]) for key in matched_keys)
    stats = {
        "left_row_count": len(left_rows),
        "right_row_count": len(right_rows),
        "left_non_null_key_count": sum(left_index.values()),
        "right_non_null_key_count": sum(len(rows) for rows in right_index.values()),
        "left_distinct_key_count": len(left_index),
        "right_distinct_key_count": len(right_index),
        "matched_distinct_keys": len(matched_keys),
        "matched_left_row_count": matched_left_rows,
        "matched_right_row_count": matched_right_rows,
        "joined_pair_count": joined_pair_count,
        "unmatched_left_row_count": len(left_rows) - matched_left_rows,
        "unmatched_right_row_count": len(right_rows) - matched_right_rows,
        "left_duplicate_key_count": sum(count > 1 for count in left_index.values()),
        "right_duplicate_key_count": sum(len(rows) > 1 for rows in right_index.values()),
        "sample_matched_keys": matched_keys[:20],
    }
    return stats, right_index


def _joined_row(left_row: dict[str, Any], right_row: dict[str, Any] | None) -> dict[str, Any]:
    """Namespace both sides so a joined handle never silently overwrites fields."""

    joined = {f"left__{key}": value for key, value in left_row.items()}
    if right_row is not None:
        joined.update({f"right__{key}": value for key, value in right_row.items()})
    return joined


def _expand_relation_graph(
    tables: dict[str, dict[str, Any]],
    steps: list[dict[str, Any]],
    table_rows: dict[str, list[dict[str, Any]]],
    row_limit: int,
) -> tuple[list[dict[str, Any]], bool]:
    """Execute a validated relation tree without exposing internal join nodes."""

    root_id = next(iter(tables))
    root_alias = tables[root_id]["alias"]
    current = [{f"{root_alias}__{key}": value for key, value in row.items()} for row in table_rows[root_id]]
    for step in steps:
        existing_alias = tables[step["existing_table_id"]]["alias"]
        new_table = tables[step["new_table_id"]]
        new_alias = new_table["alias"]
        new_columns = list(new_table["columns"])
        index: dict[str, list[dict[str, Any]]] = {}
        for row in table_rows[step["new_table_id"]]:
            key = _relation_key(row.get(step["new_key"]))
            if key is not None:
                index.setdefault(key, []).append(row)
        expanded: list[dict[str, Any]] = []
        for joined_row in current:
            key = _relation_key(joined_row.get(f"{existing_alias}__{step['existing_key']}"))
            matches = index.get(key, []) if key is not None else []
            if matches:
                for row in matches:
                    expanded.append({**joined_row, **{f"{new_alias}__{field}": value for field, value in row.items()}})
            elif step["join_type"] == "left":
                expanded.append({**joined_row, **{f"{new_alias}__{field}": None for field in new_columns}})
            if len(expanded) > row_limit:
                return expanded[: row_limit + 1], True
        current = expanded
    return current, False


def _matches(row: dict[str, Any], predicate: dict[str, Any]) -> bool:
    value, target, operator = row.get(predicate["source_column"]), predicate.get("value"), predicate["operator"]
    if operator == "between":
        number = _number(value); bounds = target if isinstance(target, list) else []
        return number is not None and len(bounds) == 2 and _number(bounds[0]) is not None and _number(bounds[1]) is not None and min(float(bounds[0]), float(bounds[1])) <= number <= max(float(bounds[0]), float(bounds[1]))
    if operator == "contains": return str(target).casefold() in str(value or "").casefold()
    values = target if isinstance(target, list) else [target]
    normalized = {str(item).casefold() for item in values}
    return str(value).casefold() in normalized


def _unique_headers(values: Iterable[Any]) -> list[str]:
    headers, counts = [], Counter()
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value not in (None, "") else f"column_{index}"
        counts[base] += 1
        headers.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return headers


def _quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'
