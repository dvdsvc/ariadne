"""Read-only adapter for Wu et al. global zircon U-Th-Pb database v2.1.

The SQLite file is only a query index over the four source workbooks.  It does
not reinterpret, aggregate, clean, or overwrite the grain-level records.
"""

from __future__ import annotations

import hashlib
import json
import sqlite3
from collections import Counter
from pathlib import Path
from typing import Any, Callable

from .contracts import CoverageProfile, SourceDescriptor
from .data_operation_contract import add_artifact_contract


DATABASE_FILES = tuple(f"Database_v2.1_part{index}.xlsx" for index in (1, 2, 3))
REFERENCE_FILE = "References_v2.1.xlsx"
NATIVE_FIELDS = (
    "Ref_number", "Data_number", "Sample_number",
    "isotope206Pb/238U", "isotope206Pb/238U_σ",
    "isotope207Pb/235U", "isotope207Pb/235U_σ",
    "isotope207Pb/206Pb", "isotope207Pb/206Pb_σ",
    "isotope208Pb/232Th", "isotope208Pb/232Th_σ",
    "age206Pb/238U", "age206Pb/238U_σ",
    "age207Pb/235U", "age207Pb/235U_σ",
    "age207Pb/206Pb", "age207Pb/206Pb_σ",
    "age208Pb/232Th", "age208Pb/232Th_σ",
    "Instrument", "Reference_material", "Lithology", "Longitude", "Latitude",
)
SQL_COLUMNS = (
    "ref_number", "data_number", "sample_number",
    "ratio_206_238", "ratio_206_238_sigma", "ratio_207_235", "ratio_207_235_sigma",
    "ratio_207_206", "ratio_207_206_sigma", "ratio_208_232", "ratio_208_232_sigma",
    "age_206_238", "age_206_238_sigma", "age_207_235", "age_207_235_sigma",
    "age_207_206", "age_207_206_sigma", "age_208_232", "age_208_232_sigma",
    "instrument", "reference_material", "lithology", "longitude", "latitude",
)
AGE_CHANNELS = {
    "206Pb/238U": ("age_206_238", "age_206_238_sigma", "age206Pb/238U", "age206Pb/238U_σ"),
    "207Pb/235U": ("age_207_235", "age_207_235_sigma", "age207Pb/235U", "age207Pb/235U_σ"),
    "207Pb/206Pb": ("age_207_206", "age_207_206_sigma", "age207Pb/206Pb", "age207Pb/206Pb_σ"),
    "208Pb/232Th": ("age_208_232", "age_208_232_sigma", "age208Pb/232Th", "age208Pb/232Th_σ"),
}


def source_signature(data_directory: str | Path) -> str:
    root = Path(data_directory)
    files = [root / name for name in (*DATABASE_FILES, REFERENCE_FILE)]
    missing = [str(path) for path in files if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Global zircon database files are missing: {missing}")
    payload = [(path.name, path.stat().st_size, path.stat().st_mtime_ns) for path in files]
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


def build_global_zircon_cache(
    data_directory: str | Path,
    cache_path: str | Path,
    *,
    progress: Callable[[str], None] | None = None,
) -> Path:
    """Stream the native workbooks into an indexed SQLite read model."""

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as error:  # pragma: no cover
        raise RuntimeError("Global zircon cache construction requires openpyxl") from error
    root, target = Path(data_directory), Path(cache_path)
    if target.exists():
        raise FileExistsError(f"Refusing to overwrite existing cache: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(target)
    try:
        connection.executescript("""
            PRAGMA journal_mode=OFF;
            PRAGMA synchronous=OFF;
            PRAGMA temp_store=MEMORY;
            CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE zircon_records (
                ref_number TEXT, data_number INTEGER PRIMARY KEY, sample_number TEXT,
                ratio_206_238 REAL, ratio_206_238_sigma REAL,
                ratio_207_235 REAL, ratio_207_235_sigma REAL,
                ratio_207_206 REAL, ratio_207_206_sigma REAL,
                ratio_208_232 REAL, ratio_208_232_sigma REAL,
                age_206_238 REAL, age_206_238_sigma REAL,
                age_207_235 REAL, age_207_235_sigma REAL,
                age_207_206 REAL, age_207_206_sigma REAL,
                age_208_232 REAL, age_208_232_sigma REAL,
                instrument TEXT, reference_material TEXT, lithology TEXT,
                longitude REAL, latitude REAL
            );
            CREATE TABLE references_data (
                ref_number TEXT PRIMARY KEY, author_surname TEXT, author_given_name TEXT,
                year_publication TEXT, title TEXT, journal TEXT, volume TEXT, issue TEXT,
                first_page TEXT, last_page TEXT, doi TEXT, link TEXT, access TEXT
            );
        """)
        insert_records = f"INSERT INTO zircon_records ({','.join(SQL_COLUMNS)}) VALUES ({','.join('?' for _ in SQL_COLUMNS)})"
        total = 0
        for filename in DATABASE_FILES:
            if progress:
                progress(f"reading {filename}")
            workbook = load_workbook(root / filename, read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            rows = sheet.iter_rows(min_row=4, values_only=True)
            headers = tuple(next(rows))
            if headers != NATIVE_FIELDS:
                raise ValueError(f"Unexpected schema in {filename}: {headers}")
            batch: list[tuple[Any, ...]] = []
            for row in rows:
                batch.append(tuple(row))
                if len(batch) >= 10_000:
                    connection.executemany(insert_records, batch)
                    total += len(batch)
                    batch.clear()
                    if progress and total % 200_000 == 0:
                        progress(f"indexed {total:,} grain analyses")
            if batch:
                connection.executemany(insert_records, batch)
                total += len(batch)
            connection.commit()
            workbook.close()
        reference_workbook = load_workbook(root / REFERENCE_FILE, read_only=True, data_only=True)
        reference_sheet = reference_workbook[reference_workbook.sheetnames[0]]
        reference_rows = reference_sheet.iter_rows(min_row=4, values_only=True)
        reference_headers = tuple(next(reference_rows))
        expected_reference_headers = (
            "Ref_number", "Author_surname", "Author_given_name", "Year_publication", "Title",
            "Journal", "Volume", "Issue", "First_page", "Last_page", "DOI", "Link", "Access",
        )
        if reference_headers != expected_reference_headers:
            raise ValueError(f"Unexpected reference schema: {reference_headers}")
        connection.executemany(
            "INSERT INTO references_data VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (tuple(row) for row in reference_rows),
        )
        reference_workbook.close()
        if progress:
            progress("building query indexes")
        connection.executescript("""
            CREATE INDEX idx_zircon_ref ON zircon_records(ref_number);
            CREATE INDEX idx_zircon_sample ON zircon_records(ref_number, sample_number);
            CREATE INDEX idx_zircon_age206238 ON zircon_records(age_206_238);
            CREATE INDEX idx_zircon_age207235 ON zircon_records(age_207_235);
            CREATE INDEX idx_zircon_age207206 ON zircon_records(age_207_206);
            CREATE INDEX idx_zircon_age208232 ON zircon_records(age_208_232);
            CREATE INDEX idx_zircon_position ON zircon_records(longitude, latitude);
            CREATE INDEX idx_zircon_lithology ON zircon_records(lithology);
            CREATE INDEX idx_zircon_instrument ON zircon_records(instrument);
        """)
        metadata = {
            "source_signature": source_signature(root), "record_count": str(total),
            "database_version": "2.1", "source_doi": "10.5281/zenodo.8303076",
            "license": "CC BY 4.0", "row_semantics": "single zircon U-Th-Pb analysis",
        }
        connection.executemany("INSERT INTO metadata(key,value) VALUES (?,?)", metadata.items())
        connection.commit()
        if progress:
            progress(f"cache complete: {total:,} grain analyses")
        return target
    except Exception:
        connection.close()
        if target.exists():
            target.unlink()
        raise
    finally:
        try:
            connection.close()
        except Exception:
            pass


class GlobalZirconSQLiteSource:
    """Native grain-analysis observation space backed by a reusable index."""

    source_id = "global_zircon_v2_1"

    def __init__(self, data_directory: str | Path, cache_path: str | Path, *, age_channel: str, max_query_rows: int = 250_000) -> None:
        if age_channel not in AGE_CHANNELS:
            raise ValueError(f"age_channel must be one of {sorted(AGE_CHANNELS)}")
        self.data_directory, self.cache_path = Path(data_directory), Path(cache_path)
        self.age_channel, self.max_query_rows = age_channel, max_query_rows
        if not self.cache_path.exists():
            raise FileNotFoundError(f"Global zircon query cache does not exist: {self.cache_path}")
        metadata = self._metadata()
        if metadata.get("source_signature") != source_signature(self.data_directory):
            raise ValueError("Global zircon cache does not match the current source workbooks")
        self.version = f"global-zircon:{metadata.get('database_version', '2.1')}:{metadata['source_signature'][:12]}"
        self._summary_cache: dict[str, dict[str, Any]] = {}

    @property
    def _age_columns(self) -> tuple[str, str, str, str]:
        return AGE_CHANNELS[self.age_channel]

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(f"file:{self.cache_path}?mode=ro", uri=True)
        connection.row_factory = sqlite3.Row
        return connection

    def _metadata(self) -> dict[str, str]:
        with sqlite3.connect(f"file:{self.cache_path}?mode=ro", uri=True) as connection:
            return dict(connection.execute("SELECT key,value FROM metadata"))

    def describe_source(self) -> SourceDescriptor:
        return SourceDescriptor(
            self.source_id, "Global zircon U-Th-Pb database v2.1", "event_clock_sensor", self.version,
            "Native single-zircon analysis records from Wu et al. (2023); no event interpretation or sample aggregation is added.",
            list(NATIVE_FIELDS), self.describe_capabilities(), self.describe_schema(),
            provenance={"doi": "10.5281/zenodo.8303076", "article_doi": "10.5194/essd-15-5171-2023", "license": "CC BY 4.0"},
        )

    def describe_scout_summary(self) -> dict[str, Any]:
        return {
            "source_id": self.source_id, "name": "Global zircon U-Th-Pb database v2.1",
            "sensor_kind": "event_clock_sensor", "version": self.version,
            "description": "Native single-zircon analysis records; no event interpretation or sample aggregation is added.",
            "variables": list(NATIVE_FIELDS), "capabilities": self.describe_capabilities(),
            "authoritative": True,
        }

    def describe_schema(self) -> dict[str, Any]:
        age_column, error_column, native_age, native_error = self._age_columns
        return {
            "row_semantics": "single zircon U-Th-Pb analysis",
            "native_fields": list(NATIVE_FIELDS),
            "reference_join": {"field": "Ref_number", "table": REFERENCE_FILE},
            "canonical_projection": {
                "age_ma": {"source_field": native_age, "sql_column": age_column, "selection": "explicit runtime configuration"},
                "age_error_ma": {"source_field": native_error, "sql_column": error_column, "uncertainty_form": "1 standard deviation"},
                "isotope_system": {"constant": "U-Pb"}, "mineral": {"constant": "zircon"},
                "rock_type": {"source_field": "Lithology"}, "technique": {"source_field": "Instrument"},
                "spatial_location": {"source_fields": ["Longitude", "Latitude"], "uncertainty": "not supplied per record"},
            },
            "claim_limits": [
                "Rows are analyses, not independently interpreted geological events.",
                "The source supplies four age systems; canonical age_ma uses only the explicitly configured channel.",
                "Coordinates have no per-record uncertainty or resolution field.",
            ],
        }

    def describe_capabilities(self) -> list[str]:
        return ["event_age", "age_uncertainty", "source_provenance", "spatial_location", "mineral"]

    def describe_observables(self, scope: dict[str, Any]) -> dict[str, dict[str, Any]]:
        summary = self._summary(scope)
        count = summary["record_count"]
        def item(present: int, fields: list[str], **extra: Any) -> dict[str, Any]:
            return {"available": present > 0, "coverage": round(present / count, 4) if count else 0.0, "record_count": present, "source_fields": fields, **extra}
        return {
            "event_age": item(summary["age_count"], [self._age_columns[2]], age_channel=self.age_channel),
            "age_uncertainty": item(summary["age_error_count"], [self._age_columns[3]], uncertainty_form="1 standard deviation"),
            "source_provenance": item(summary["reference_count"], ["Ref_number"], reference_join_available=True),
            "spatial_location": item(summary["coordinate_count"], ["Longitude", "Latitude"]),
            "spatial_uncertainty": item(0, [], limitation="not supplied per record"),
            "event_interpretation": item(0, [], limitation="not a field in this grain-analysis database"),
            "mineral": item(count, [], constant="zircon"),
            "multi_clock": item(0, [], clock_systems=["U-Pb"], system_count=1),
            "material_state": item(0, []),
        }

    def audit_semantic_bindings(self) -> dict[str, Any]:
        bindings = self.describe_schema()["canonical_projection"]
        return {
            "manifest_id": f"global_zircon_v2_1/{self.age_channel}", "schema_fingerprint": self.version,
            "status": "valid", "bindings": [
                {"concept_id": concept, "source_field": binding.get("source_field") or binding.get("source_fields") or binding.get("constant"), "status": "valid", "meaning": binding}
                for concept, binding in bindings.items()
            ],
        }

    def describe_query_contract(self) -> dict[str, Any]:
        return {
            "query_contract": "observation_query/v1", "source_id": self.source_id, "version": self.version,
            "fields": [
                {"field": "record_id", "semantic_role": "single_analysis_identifier", "operators": ["eq", "in"], "value_type": "string"},
                {"field": "sample_id", "semantic_role": "sample_identifier", "operators": ["eq", "in"], "value_type": "string"},
                {"field": "source_id", "semantic_role": "reference_identifier", "operators": ["eq", "in"], "value_type": "string"},
                {"field": "age_ma", "semantic_role": "configured_u_pb_age_channel", "operators": ["between"], "value_type": "number", "age_channel": self.age_channel},
                {"field": "isotope_system", "semantic_role": "clock_system", "operators": ["eq", "in"], "value_type": "string", "known_values": ["U-Pb"]},
                {"field": "rock_type", "semantic_role": "host_rock_class", "operators": ["eq", "in"], "value_type": "string", "known_values": self._distinct("lithology")},
                {"field": "technique", "semantic_role": "dating_instrument", "operators": ["eq", "in"], "value_type": "string", "known_values": self._distinct("instrument")},
                {"field": "longitude", "semantic_role": "present_day_coordinate", "operators": ["between"], "value_type": "number"},
                {"field": "latitude", "semantic_role": "present_day_coordinate", "operators": ["between"], "value_type": "number"},
            ],
            "failure_mode": "Unsupported or partially resolved scopes return zero rows. Queries exceeding max_query_rows fail instead of truncating.",
            "max_query_rows": self.max_query_rows,
        }

    def workspace_catalogue(self) -> dict[str, Any]:
        """Live relational surface for ``data.tool`` actions.

        These are canonical fields deliberately exposed by this native adapter,
        not a guessed mapping over spreadsheet headers.  The SQLite index owns
        how each field is translated to its physical column.
        """

        fields = ["record_id", "sample_id", "source_id", "age_ma", "age_error_ma", "isotope_system", "mineral", "rock_type", "technique", "latitude", "longitude"]
        return {
            "source_id": self.source_id,
            "data_operation_contract": "data_operation/v1",
            "decision_contract": "data_decision_required/v1",
            "tables": [{"table_id": "zircon_records", "fields": fields}],
            "operations": ["inspect", "profile", "sample", "query", "aggregate", "materialize"],
            "aggregate_operations": ["count", "min", "max", "mean", "n_unique"],
            "filter_operators": ["eq", "in", "between"],
            "limitations": ["Rows are single zircon analyses, not interpreted geological events.", "Median and standard deviation aggregation are not exposed through SQLite because this adapter does not emulate unsupported database functions."],
        }

    def workspace_execute(self, request: dict[str, Any]) -> dict[str, Any]:
        """Execute a validated relational operation inside the SQLite index."""

        operation = str(request.get("operation") or "")
        if operation == "inspect":
            return {"summary": "Inspected the native global-zircon relational surface.", "findings": {"catalogue": self.workspace_catalogue(), "schema": self.describe_schema(), "query_contract": self.describe_query_contract()}}
        scope = self._workspace_scope(request)
        if operation == "profile":
            summary = self._summary(scope)
            return {"summary": f"Profiled {summary['record_count']:,} zircon analyses in the requested scope.", "findings": summary, "limitations": list(self.describe_schema()["claim_limits"])}
        if operation in {"sample", "query"}:
            count = self._summary(scope)["record_count"]
            limit = max(1, min(int(request.get("row_limit") or 100), 2_000))
            rows = self._select(scope, limit=limit, enforce_limit=False)
            return {"summary": f"Read {len(rows):,} of {count:,} matching zircon analyses.", "findings": {"record_count": count, "returned_records": len(rows), "sample": rows[:20], "truncated": count > len(rows)}}
        if operation == "materialize":
            count = self._summary(scope)["record_count"]
            columns = list(dict.fromkeys(str(item) for item in request.get("columns", []) if item))
            if not columns:
                columns = ["record_id", "sample_id", "source_id", "age_ma", "age_error_ma", "rock_type", "technique", "latitude", "longitude"]
            handle = add_artifact_contract(
                {"handle_contract": "data_handle/v1", "source_id": self.source_id, "scope": scope, "columns": columns, "dataset_name": str(request.get("dataset_name") or "")},
                source_id=self.source_id, operation="materialize", request=request, record_count=count,
            )
            return {
                "summary": f"Created a durable query handle for {count:,} matching zircon analyses; rows were not placed in model context.",
                "findings": {"record_count": count, "dataset_handle": handle},
                "limitations": ["The handle is a reproducible source query specification. A numerical provider must explicitly consume it; it is not a Python row list."],
            }
        if operation == "aggregate":
            return self._workspace_aggregate(scope, request)
        raise ValueError(f"Unsupported global-zircon workspace operation: {operation!r}")

    def load_workspace_handle(self, handle: dict[str, Any]) -> list[dict[str, Any]]:
        """Load just the handle's declared canonical fields for a local provider.

        This deliberately bypasses ``max_query_rows``: the row limit protects
        model-facing raw queries, whereas a materialized handle is a local
        numerical job and never enters model context.
        """

        if handle.get("source_id") != self.source_id:
            raise ValueError("Workspace handle belongs to a different source")
        requested = list(dict.fromkeys(str(item) for item in handle.get("columns", []) if item))
        allowed = {"record_id", "sample_id", "source_id", "age_ma", "age_error_ma", "isotope_system", "mineral", "rock_type", "technique", "latitude", "longitude"}
        if not requested or not set(requested).issubset(allowed):
            raise ValueError("Workspace handle has invalid global-zircon columns")
        scope = handle.get("scope")
        if not isinstance(scope, dict):
            raise ValueError("Workspace handle lacks its source scope")
        where, parameters = self._where(scope)
        expressions = {
            "record_id": "data_number", "sample_id": "sample_number", "source_id": "ref_number",
            "age_ma": self._age_columns[0], "age_error_ma": self._age_columns[1],
            "rock_type": "lithology", "technique": "instrument", "latitude": "latitude", "longitude": "longitude",
        }
        select = []
        for field in requested:
            if field == "isotope_system": select.append("'U-Pb' AS isotope_system")
            elif field == "mineral": select.append("'zircon' AS mineral")
            else: select.append(f"{expressions[field]} AS {field}")
        with self._connect() as connection:
            rows = [dict(row) for row in connection.execute(f"SELECT {', '.join(select)} FROM zircon_records WHERE {where} ORDER BY data_number", parameters)]
        for row in rows:
            row["record_id"] = f"zircon:{row['record_id']}" if row.get("record_id") is not None else None
            row["position"] = {"latitude": row.get("latitude"), "longitude": row.get("longitude"), "uncertainty_km": None, "resolution": "unresolved", "provenance": "database_compiled_coordinate; per-record uncertainty not supplied"}
        return rows

    def _workspace_scope(self, request: dict[str, Any]) -> dict[str, Any]:
        valid = {item["field"] for item in self.describe_query_contract()["fields"]}
        predicates = []
        for raw in request.get("filters", []):
            field, operator = str(raw.get("field") or ""), str(raw.get("operator") or "")
            if field not in valid:
                raise ValueError(f"Unknown global-zircon workspace field: {field!r}")
            if operator == "between":
                value = [raw.get("minimum_number"), raw.get("maximum_number")]
            elif operator in {"eq", "in"}:
                value = list(raw.get("text_values") or [])
            else:
                raise ValueError(f"Unsupported global-zircon workspace filter: {field!r}/{operator!r}")
            predicates.append({"field": field, "operator": operator, "value": value})
        return {"_scope_contract": "observation_scope/v1", "source_id": self.source_id, "status": "resolved", "predicates": predicates}

    def _workspace_expression(self, field: str) -> str:
        age, error, _, _ = self._age_columns
        expressions = {
            "record_id": "data_number", "sample_id": "sample_number", "source_id": "ref_number",
            "age_ma": age, "age_error_ma": error, "rock_type": "lithology", "technique": "instrument",
            "latitude": "latitude", "longitude": "longitude", "isotope_system": "'U-Pb'", "mineral": "'zircon'",
        }
        if field not in expressions:
            raise ValueError(f"Unknown global-zircon workspace field: {field!r}")
        return expressions[field]

    def _workspace_aggregate(self, scope: dict[str, Any], request: dict[str, Any]) -> dict[str, Any]:
        where, where_parameters = self._where(scope)
        group_by = [str(field) for field in request.get("group_by", []) if field]
        bin_field = str(request.get("bin_field") or "")
        bin_width = float(request.get("bin_width") or 0)
        bin_origin = float(request.get("bin_origin") or 0)
        select_parts: list[str] = []
        parameters: list[Any] = []
        aliases: list[str] = []
        for field in group_by:
            alias = f"group_{len(aliases)}"
            select_parts.append(f"{self._workspace_expression(field)} AS {alias}")
            aliases.append(alias)
        if bin_field:
            if bin_width <= 0:
                raise ValueError("A workspace age bin requires a positive bin_width")
            alias = "bin_start"
            expression = self._workspace_expression(bin_field)
            select_parts.append(f"(? + CAST(({expression} - ?) / ? AS INTEGER) * ?) AS {alias}")
            parameters.extend([bin_origin, bin_origin, bin_width, bin_width])
            aliases.append(alias)
        for raw in request.get("aggregates", []) or [{"field": "", "operation": "count", "as": "count"}]:
            operation, field = str(raw.get("operation") or "count"), str(raw.get("field") or "")
            alias = str(raw.get("as") or f"{operation}_{field or 'records'}")
            if operation == "count":
                expression = self._workspace_expression(field) if field else "*"
                select_parts.append(f"COUNT({expression}) AS {alias}")
            elif operation == "n_unique":
                select_parts.append(f"COUNT(DISTINCT {self._workspace_expression(field)}) AS {alias}")
            elif operation in {"min", "max", "mean"}:
                function = {"min": "MIN", "max": "MAX", "mean": "AVG"}[operation]
                select_parts.append(f"{function}({self._workspace_expression(field)}) AS {alias}")
            else:
                raise ValueError(f"Global-zircon SQLite aggregation does not expose {operation!r}")
        sql = f"SELECT {', '.join(select_parts)} FROM zircon_records WHERE {where}"
        if aliases:
            sql += " GROUP BY " + ", ".join(aliases)
            sql += " ORDER BY " + ", ".join(aliases)
        limit = max(1, min(int(request.get("row_limit") or 500), 10_000))
        sql += " LIMIT ?"
        parameters.extend(where_parameters); parameters.append(limit + 1)
        with self._connect() as connection:
            rows = [dict(row) for row in connection.execute(sql, parameters).fetchall()]
        truncated = len(rows) > limit
        return {
            "summary": f"Aggregated {self._summary(scope)['record_count']:,} zircon analyses into {min(len(rows), limit):,} group(s) inside SQLite.",
            "findings": {"input_record_count": self._summary(scope)["record_count"], "group_count_returned": min(len(rows), limit), "rows": rows[:limit], "truncated": truncated},
            "limitations": list(self.workspace_catalogue()["limitations"]),
        }

    def describe_discovery_coverage(self) -> dict[str, Any]:
        summary = self._summary({})
        return {
            "record_count": summary["record_count"], "row_semantics": "single zircon analysis",
            "time_coverage_ma": {"min": summary["minimum_age_ma"], "max": summary["maximum_age_ma"], "age_channel": self.age_channel},
            "named_spatial_coverage": {}, "coordinate_coverage": summary["coordinate_coverage"],
            "coordinate_bbox": summary["coordinate_bbox"],
            "lithologies": self._counts("lithology"), "instruments": self._counts("instrument"),
            "reference_count": self._reference_total(),
        }

    def profile_coverage(self, scope: dict[str, Any]) -> CoverageProfile:
        summary, total = self._summary(scope), self._summary({})["record_count"]
        count = summary["record_count"]
        missingness = {
            "age_ma": self._missing(count, summary["age_count"]), "age_error_ma": self._missing(count, summary["age_error_count"]),
            "latitude": self._missing(count, summary["coordinate_count"]), "longitude": self._missing(count, summary["coordinate_count"]),
            "rock_type": self._missing(count, summary["lithology_count"]), "source_id": self._missing(count, summary["reference_count"]),
            "event_interpretation": 1.0, "mineral": 0.0 if count else 1.0, "technique": self._missing(count, summary["instrument_count"]),
        }
        return CoverageProfile(
            self.source_id, count, total,
            {"fields": list(NATIVE_FIELDS), "isotope_systems": {"U-Pb": count}, "row_semantics": "single zircon analysis", "age_channel": self.age_channel, "observables": self.describe_observables(scope)},
            missingness, self.describe_resolution(scope), self.describe_uncertainty(scope), self.estimate_query(scope), self.sample_records(scope), self._warnings(summary), self.version,
        )

    def estimate_query(self, scope: dict[str, Any]) -> dict[str, Any]:
        count = self._summary(scope)["record_count"]
        return {"estimated_records": count, "estimated_cost": "low" if count < 5_000 else ("moderate" if count < 100_000 else "high"), "executable": count <= self.max_query_rows, "query_scope": scope}

    def sample_records(self, scope: dict[str, Any], n: int = 5) -> list[dict[str, Any]]:
        return self._select(scope, limit=max(0, n), enforce_limit=False)

    def describe_resolution(self, scope: dict[str, Any]) -> dict[str, Any]:
        summary = self._summary(scope)
        return {
            "coordinate_coverage": summary["coordinate_coverage"],
            "position_resolution_counts": {"unresolved_uncertainty": summary["coordinate_count"]},
            "median_spatial_uncertainty_km": None, "minimum_resolvable_migration_km": None,
            "age_resolution": self.age_channel if summary["age_count"] else "unavailable",
            "limitation": "Coordinates are available but per-record spatial uncertainty and resolution are absent.",
        }

    def describe_uncertainty(self, scope: dict[str, Any]) -> dict[str, Any]:
        summary = self._summary(scope)
        return {
            "age_error_coverage": round(summary["age_error_count"] / summary["record_count"], 4) if summary["record_count"] else 0.0,
            "median_age_error_ma": summary["median_age_error_ma"], "uncertainty_form": "1 standard deviation",
            "median_spatial_uncertainty_km": None, "spatial_uncertainty_available": False,
        }

    def query(self, scope: dict[str, Any]) -> list[dict[str, Any]]:
        count = self._summary(scope)["record_count"]
        if count > self.max_query_rows:
            raise ValueError(f"Global zircon query would return {count:,} rows; narrow the scope below {self.max_query_rows:,} rows")
        return self._select(scope, limit=None, enforce_limit=True)

    def _select(self, scope: dict[str, Any], *, limit: int | None, enforce_limit: bool) -> list[dict[str, Any]]:
        where, parameters = self._where(scope)
        age_column, error_column, _, _ = self._age_columns
        sql = f"""
            SELECT z.*, r.author_surname, r.year_publication, r.title, r.journal, r.doi, r.link
            FROM zircon_records z LEFT JOIN references_data r USING(ref_number)
            WHERE {where} ORDER BY z.data_number
        """
        if limit is not None:
            sql += " LIMIT ?"
            parameters = [*parameters, limit]
        with self._connect() as connection:
            rows = connection.execute(sql, parameters).fetchall()
        return [self._project(dict(row), age_column, error_column) for row in rows]

    def _project(self, row: dict[str, Any], age_column: str, error_column: str) -> dict[str, Any]:
        position = {
            "latitude": row.get("latitude"), "longitude": row.get("longitude"),
            "uncertainty_km": None, "resolution": "unresolved",
            "provenance": "database_compiled_coordinate; per-record uncertainty not supplied",
        }
        native = {native: row.get(sql_name) for native, sql_name in zip(NATIVE_FIELDS, SQL_COLUMNS)}
        return {
            "record_id": f"zircon:{row['data_number']}", "data_number": row["data_number"],
            "paper_id": row.get("ref_number"), "source_id": row.get("ref_number"), "sample_id": row.get("sample_number"),
            "age_ma": row.get(age_column), "age_error_ma": row.get(error_column), "age_channel": self.age_channel,
            "isotope_system": "U-Pb", "mineral": "zircon", "technique": row.get("instrument"), "rock_type": row.get("lithology"),
            "latitude": row.get("latitude"), "longitude": row.get("longitude"), "position": position,
            "reference_payload": {key: row.get(key) for key in ("author_surname", "year_publication", "title", "journal", "doi", "link")},
            "native_payload": native,
            "_semantic_field_provenance": {
                "event_age": [self._age_columns[2]], "age_uncertainty": [self._age_columns[3]],
                "source_provenance": ["Ref_number"], "spatial_location": ["Longitude", "Latitude"],
                "mineral": ["constant:zircon"],
            },
        }

    def _where(self, scope: dict[str, Any]) -> tuple[str, list[Any]]:
        if scope.get("_scope_contract") == "observation_scope/v1" and scope.get("status") in {"partial", "scope_unresolvable"}:
            return "0", []
        predicates = scope.get("predicates", []) if scope.get("_scope_contract") == "observation_scope/v1" else []
        clauses: list[str] = []
        parameters: list[Any] = []
        field_map = {
            "record_id": "data_number", "sample_id": "sample_number", "source_id": "ref_number",
            "age_ma": self._age_columns[0], "rock_type": "lithology", "technique": "instrument",
            "longitude": "longitude", "latitude": "latitude",
        }
        for predicate in predicates:
            field, operator, value = predicate.get("field"), predicate.get("operator"), predicate.get("value")
            if field == "isotope_system":
                values = value if isinstance(value, list) else [value]
                if any(str(item).casefold() != "u-pb".casefold() for item in values):
                    return "0", []
                continue
            column = field_map.get(field)
            if column is None:
                raise ValueError(f"Unsupported global zircon predicate: {field!r}")
            if operator == "between" and isinstance(value, list) and len(value) == 2:
                clauses.append(f"{column} BETWEEN ? AND ?")
                parameters.extend([min(float(value[0]), float(value[1])), max(float(value[0]), float(value[1]))])
            elif operator in {"eq", "in"}:
                values = value if isinstance(value, list) else [value]
                clauses.append(f"LOWER({column}) IN ({','.join('?' for _ in values)})")
                parameters.extend(str(item).casefold() for item in values)
            elif operator == "contains":
                values = value if isinstance(value, list) else [value]
                non_empty = [str(item).casefold() for item in values if str(item)]
                if not non_empty:
                    return "0", []
                clauses.append("(" + " OR ".join(f"LOWER({column}) LIKE ?" for _ in non_empty) + ")")
                parameters.extend(f"%{item}%" for item in non_empty)
            else:
                raise ValueError(f"Unsupported global zircon predicate: {field!r}/{operator!r}")
        return " AND ".join(clauses) if clauses else "1", parameters

    def _summary(self, scope: dict[str, Any]) -> dict[str, Any]:
        key = json.dumps(scope, sort_keys=True, default=str)
        if key in self._summary_cache:
            return self._summary_cache[key]
        where, parameters = self._where(scope)
        age, error, _, _ = self._age_columns
        sql = f"""
            SELECT COUNT(*) record_count,
                   COUNT({age}) age_count, COUNT({error}) age_error_count,
                   COUNT(ref_number) reference_count, COUNT(instrument) instrument_count,
                   COUNT(lithology) lithology_count,
                   SUM(CASE WHEN longitude IS NOT NULL AND latitude IS NOT NULL THEN 1 ELSE 0 END) coordinate_count,
                   MIN({age}) minimum_age_ma, MAX({age}) maximum_age_ma,
                   MIN(longitude) minimum_longitude, MAX(longitude) maximum_longitude,
                   MIN(latitude) minimum_latitude, MAX(latitude) maximum_latitude
            FROM zircon_records WHERE {where}
        """
        with self._connect() as connection:
            result = dict(connection.execute(sql, parameters).fetchone())
        count, coordinate_count = int(result["record_count"] or 0), int(result["coordinate_count"] or 0)
        result.update({
            "record_count": count, "age_count": int(result["age_count"] or 0), "age_error_count": int(result["age_error_count"] or 0),
            "reference_count": int(result["reference_count"] or 0), "instrument_count": int(result["instrument_count"] or 0), "lithology_count": int(result["lithology_count"] or 0),
            "coordinate_count": coordinate_count, "coordinate_coverage": round(coordinate_count / count, 4) if count else 0.0,
            "coordinate_bbox": {"min_longitude": result["minimum_longitude"], "max_longitude": result["maximum_longitude"], "min_latitude": result["minimum_latitude"], "max_latitude": result["maximum_latitude"]} if coordinate_count else None,
            "median_age_error_ma": None,
        })
        self._summary_cache[key] = result
        return result

    def _distinct(self, column: str) -> list[str]:
        with self._connect() as connection:
            return [row[0] for row in connection.execute(f"SELECT DISTINCT {column} FROM zircon_records WHERE {column} IS NOT NULL ORDER BY {column}")]

    def _counts(self, column: str) -> dict[str, int]:
        with self._connect() as connection:
            return {str(row[0]): int(row[1]) for row in connection.execute(f"SELECT {column}, COUNT(*) FROM zircon_records WHERE {column} IS NOT NULL GROUP BY {column} ORDER BY COUNT(*) DESC")}

    def _reference_total(self) -> int:
        with self._connect() as connection:
            return int(connection.execute("SELECT COUNT(*) FROM references_data").fetchone()[0])

    @staticmethod
    def _missing(total: int, present: int) -> float:
        return round(1 - present / total, 4) if total else 1.0

    def _warnings(self, summary: dict[str, Any]) -> list[str]:
        warnings = [f"Canonical age_ma is explicitly projected from {self.age_channel}; other native age channels remain in native_payload."]
        if summary["record_count"] > self.max_query_rows:
            warnings.append("This scope is too broad for row materialization; narrow it before a Researcher source query.")
        if summary["coordinate_count"]:
            warnings.append("Coordinates lack per-record uncertainty; uncertainty-dependent spatial claims require an additional constraint or must remain a gap.")
        warnings.append("Rows are grain analyses and contain no author event interpretation.")
        return warnings
