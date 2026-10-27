"""Read-only adapters for existing observation databases, never document extraction."""

from __future__ import annotations

import json
import hashlib
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from statistics import median
from typing import Any

from .contracts import CoverageProfile, SensorKind, SourceDescriptor
from .evidence_contract import observable_manifest_from_rows


@dataclass(frozen=True)
class SemanticBinding:
    """One exact, source-owned binding from a physical field to a concept."""

    concept_id: str
    source_field: str
    transform: str = "identity"
    meaning: str = ""
    claim_limit: str = ""


CANONICAL_IDENTITY_BINDINGS = tuple(
    SemanticBinding(field, field)
    for field in (
        "age_ma", "age_error_ma", "latitude", "longitude", "region", "province",
        "city", "basin", "locality", "rock_type", "source_id", "event_interpretation",
        "isotope_system", "mineral", "technique", "statistical_model",
    )
)


def load_semantic_manifest(path: str | Path) -> tuple[str, tuple[SemanticBinding, ...]]:
    """Load a source-owned semantic catalog; reject malformed bindings early."""

    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    manifest_id = payload.get("manifest_id")
    bindings = payload.get("bindings")
    if not isinstance(manifest_id, str) or not manifest_id or not isinstance(bindings, list):
        raise ValueError("semantic manifest requires manifest_id and bindings")
    try:
        resolved = tuple(SemanticBinding(**item) for item in bindings)
    except (TypeError, ValueError) as error:
        raise ValueError(f"semantic manifest has invalid binding: {error}") from error
    if not resolved or len({item.concept_id for item in resolved}) != len(resolved):
        raise ValueError("semantic manifest requires unique non-empty concept bindings")
    return manifest_id, resolved


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_isotope_system(value: Any) -> str:
    text = str(value or "U-Pb").lower().replace(" ", "")
    if "u-th-pb" in text or "u-pb" in text or "upb" in text:
        return "U-Pb"
    if "ar-ar" in text or "40ar/39ar" in text or "40ar-39ar" in text:
        return "Ar-Ar"
    if "k-ar" in text or "kar" in text:
        return "K-Ar"
    return str(value)


def normalize_record(raw: dict[str, Any], *, record_id: str, bindings: tuple[SemanticBinding, ...]) -> dict[str, Any]:
    row = dict(raw)
    normalized: dict[str, Any] = {"record_id": record_id}
    provenance: dict[str, list[str]] = {}
    for binding in bindings:
        value = row.get(binding.source_field)
        if value not in (None, ""):
            if binding.transform == "number":
                value = _number(value)
            elif binding.transform == "isotope_system":
                value = normalize_isotope_system(value)
            normalized[binding.concept_id] = value
            provenance.setdefault(binding.concept_id, []).append(binding.source_field)
    normalized["_semantic_field_provenance"] = provenance
    for key in ("age_ma", "age_error_ma", "latitude", "longitude"):
        normalized[key] = _number(normalized.get(key))
    normalized["isotope_system"] = normalize_isotope_system(normalized.get("isotope_system"))
    if normalized["latitude"] is not None and normalized["longitude"] is not None:
        normalized["position"] = {
            "latitude": normalized["latitude"], "longitude": normalized["longitude"],
            "uncertainty_km": _number(row.get("position_uncertainty_km")) or 10.0,
            "resolution": row.get("position_resolution", "locality"),
            "provenance": row.get("position_provenance", "unspecified_coordinate"),
        }
    else:
        normalized["position"] = {"latitude": None, "longitude": None, "uncertainty_km": None, "resolution": "unresolved", "provenance": "unresolved"}
    normalized.update({key: value for key, value in row.items() if key not in normalized})
    return normalized


def _in_scope(row: dict[str, Any], scope: dict[str, Any]) -> bool:
    if scope.get("_scope_contract") == "observation_scope/v1":
        return _matches_canonical_scope(row, scope)
    unknown = {key for key in scope if not key.startswith("_")} - {
        "min_age_ma", "max_age_ma", "age_range_ma", "isotope_system",
        "region", "regions", "province", "provinces", "basin", "basins",
        "locality", "localities", "city", "cities",
    }
    if unknown:
        raise ValueError(f"Uncompiled observation scope contains unsupported keys: {sorted(unknown)}")
    age = row.get("age_ma")
    min_age, max_age = _age_bounds(scope)
    if min_age is not None and (age is None or age < min_age):
        return False
    if max_age is not None and (age is None or age > max_age):
        return False
    terms = _location_terms(scope)
    if terms and not any(term in str(row.get(key, "")).lower() for term in terms for key in ("region", "locality", "province", "basin", "city")):
        return False
    system = scope.get("isotope_system")
    if system and normalize_isotope_system(system) != row.get("isotope_system"):
        return False
    return True


def _matches_canonical_scope(row: dict[str, Any], scope: dict[str, Any]) -> bool:
    if scope.get("status") in {"scope_unresolvable", "partial"}:
        return False
    predicates = scope.get("predicates", [])
    if not isinstance(predicates, list):
        raise ValueError("Canonical observation scope predicates must be a list")
    for predicate in predicates:
        if not isinstance(predicate, dict):
            raise ValueError("Canonical observation scope predicate must be an object")
        field, operator, value = predicate.get("field"), predicate.get("operator"), predicate.get("value")
        if field == "age_ma" and operator == "between":
            if not isinstance(value, list) or len(value) != 2:
                raise ValueError("age_ma between predicate requires [min, max]")
            lower, upper = _number(value[0]), _number(value[1])
            age = row.get("age_ma")
            if lower is None or upper is None or age is None or age < min(lower, upper) or age > max(lower, upper):
                return False
        elif field == "isotope_system" and operator in {"eq", "in"}:
            values = value if isinstance(value, list) else [value]
            normalized = {normalize_isotope_system(item) for item in values}
            if row.get("isotope_system") not in normalized:
                return False
        elif field in {"mineral", "rock_type"} and operator in {"eq", "in"}:
            values = value if isinstance(value, list) else [value]
            allowed = {str(item).strip().casefold() for item in values}
            if str(row.get(field, "")).strip().casefold() not in allowed:
                return False
        elif field == "location_text" and operator == "contains_any":
            values = value if isinstance(value, list) else [value]
            terms = [str(item).strip().casefold() for item in values if str(item).strip()]
            if not terms:
                raise ValueError("location_text contains_any requires at least one term")
            haystack = " ".join(str(row.get(key, "")) for key in ("region", "locality", "province", "basin", "city")).casefold()
            if not any(term in haystack for term in terms):
                return False
        else:
            raise ValueError(f"Unsupported canonical predicate: {field!r} {operator!r}")
    return True


def _scope_from_workspace_request(request: dict[str, Any], source_id: str) -> dict[str, Any]:
    """Convert the typed fallback workspace predicates to source scope.

    This is a compatibility reader for handles created before ``scope`` was
    persisted.  It only accepts the workspace predicate vocabulary; it does
    not infer synonyms or reinterpret scientific language.
    """

    predicates: list[dict[str, Any]] = []
    for predicate in request.get("filters", []) if isinstance(request.get("filters", []), list) else []:
        field = str(predicate.get("field") or "")
        operator = str(predicate.get("operator") or "")
        if field == "age_ma" and operator == "between":
            predicates.append({"field": field, "operator": operator, "value": [predicate.get("minimum_number"), predicate.get("maximum_number")]})
        elif field in {"isotope_system", "rock_type"} and operator in {"eq", "in"}:
            values = list(predicate.get("text_values") or [])
            predicates.append({"field": field, "operator": operator, "value": values[0] if operator == "eq" and values else values})
        elif field == "location_text" and operator == "contains_any":
            predicates.append({"field": field, "operator": operator, "value": list(predicate.get("text_values") or [])})
        else:
            raise ValueError(f"Workspace handle contains unsupported legacy predicate: {field}/{operator}")
    return {
        "_scope_contract": "observation_scope/v1", "source_id": source_id,
        "status": "resolved", "predicates": predicates,
    }


def _age_bounds(scope: dict[str, Any]) -> tuple[float | None, float | None]:
    """Accept the canonical bounds and the Planner's interval form."""

    lower, upper = _number(scope.get("min_age_ma")), _number(scope.get("max_age_ma"))
    interval = scope.get("age_range_ma")
    if isinstance(interval, (list, tuple)) and len(interval) == 2:
        first, second = _number(interval[0]), _number(interval[1])
        if first is not None and second is not None:
            lower = min(first, second) if lower is None else lower
            upper = max(first, second) if upper is None else upper
    return lower, upper


def _location_terms(scope: dict[str, Any]) -> tuple[str, ...]:
    """Collect caller-supplied geographic values without alias inference."""

    values: list[Any] = []
    for field in ("region", "regions", "province", "provinces", "basin", "basins", "locality", "localities", "city", "cities"):
        value = scope.get(field)
        values.extend(value if isinstance(value, (list, tuple, set)) else [value])
    terms: list[str] = []
    for value in values:
        if value in (None, ""):
            continue
        text = str(value).strip().lower()
        terms.append(text)
    return tuple(dict.fromkeys(terms))


class InMemoryObservationSource:
    """A generic sensor backed by already normalized or row-like observations."""

    def __init__(self, records: list[dict[str, Any]], *, source_id: str = "geochronology_memory", version: str = "memory", sensor_kind: SensorKind = "event_clock_sensor", name: str | None = None, capabilities: list[str] | None = None, description: str | None = None, semantic_bindings: tuple[SemanticBinding, ...] | None = None, semantic_manifest_id: str = "canonical_identity/v1") -> None:
        self.source_id, self.version, self.sensor_kind = source_id, version, sensor_kind
        self.name = name or source_id
        self.description = description or "Read-only normalized observations for scientific research."
        self.semantic_bindings = tuple(semantic_bindings or CANONICAL_IDENTITY_BINDINGS)
        self.semantic_manifest_id = semantic_manifest_id
        self.records = [normalize_record(row, record_id=str(row.get("record_id") or index + 1), bindings=self.semantic_bindings) for index, row in enumerate(records)]
        self._capabilities = capabilities or self._default_capabilities()

    def _default_capabilities(self) -> list[str]:
        if self.sensor_kind == "event_clock_sensor":
            return [term for term, description in self.describe_observables({}).items() if description["available"]]
        if self.sensor_kind == "spatial_constraint":
            return ["spatial_location", "spatial_uncertainty", "geological_polygon"]
        if self.sensor_kind == "material_state_sensor":
            return ["material_state", "source_provenance"]
        return ["source_provenance"]

    def describe_source(self) -> SourceDescriptor:
        fields = sorted({key for record in self.records for key in record})
        return SourceDescriptor(self.source_id, self.name, self.sensor_kind, self.version, self.description, fields, self._capabilities, self.describe_schema())

    def describe_scout_summary(self) -> dict[str, Any]:
        """Schema-free source identity used only as Planner context."""

        return {
            "source_id": self.source_id, "name": self.name, "sensor_kind": self.sensor_kind,
            "version": self.version, "description": self.description,
            "variables": sorted({key for record in self.records for key in record if not key.startswith("_")}),
            "capabilities": list(self._capabilities), "authoritative": True,
        }

    def describe_schema(self) -> dict[str, Any]:
        fields = sorted({key for record in self.records for key in record})
        return {
            "physical_fields": {field: {"type": "number" if field.endswith("_ma") or field in {"latitude", "longitude"} else "unknown"} for field in fields},
            "semantic_manifest": {"manifest_id": self.semantic_manifest_id, "bindings": [binding.__dict__ for binding in self.semantic_bindings]},
        }

    def describe_capabilities(self) -> list[str]:
        actual = {term for term, description in self.describe_observables({}).items() if description["available"]}
        return sorted(actual.union(self._capabilities if self.sensor_kind != "event_clock_sensor" else set()))

    def describe_observables(self, scope: dict[str, Any]) -> dict[str, dict[str, Any]]:
        rows = self.query(scope)
        provenance: dict[str, set[str]] = defaultdict(set)
        for row in rows:
            raw = row.get("_semantic_field_provenance", {})
            if not isinstance(raw, dict):
                continue
            for term, fields in raw.items():
                if isinstance(fields, list):
                    provenance[term].update(str(field) for field in fields)
        manifest = observable_manifest_from_rows(rows, field_provenance={term: sorted(fields) for term, fields in provenance.items()})
        # A material-state source contributes state through a later join, not
        # through an event-clock row returned by this source query.
        if self.sensor_kind == "material_state_sensor" and rows:
            manifest["material_state"] = {"available": True, "coverage": 1.0, "record_count": len(rows), "source_fields": ["material_state_payload"]}
        return manifest

    def audit_semantic_bindings(self) -> dict[str, Any]:
        """Verify manifest paths and current coverage against the live source.

        A binding that names a missing physical field is surfaced as broken;
        it is never silently repaired with a synonym search.
        """

        physical_fields = sorted({field for row in self.records for field in row})
        checks = []
        for binding in self.semantic_bindings:
            exists = binding.source_field in physical_fields
            present = sum(1 for row in self.records if row.get(binding.source_field) not in (None, "")) if exists else 0
            checks.append({
                "concept_id": binding.concept_id,
                "source_field": binding.source_field,
                "status": "valid" if exists else "broken",
                "coverage": round(present / len(self.records), 4) if self.records else 0.0,
                "meaning": binding.meaning,
                "claim_limit": binding.claim_limit,
            })
        fingerprint = hashlib.sha256(json.dumps({"fields": physical_fields, "bindings": [binding.__dict__ for binding in self.semantic_bindings]}, sort_keys=True).encode()).hexdigest()[:16]
        return {
            "manifest_id": self.semantic_manifest_id,
            "schema_fingerprint": fingerprint,
            "status": "valid" if all(check["status"] == "valid" for check in checks) else "binding_broken",
            "bindings": checks,
        }

    def describe_query_contract(self) -> dict[str, Any]:
        """Declare the executable semantic query language for this source.

        This is a database capability contract, not a prompt listing.  A Scope
        Compiler can only emit these field/operator combinations, and values
        shown for categorical dimensions come from this source's current rows.
        """

        def values(field: str, limit: int = 80) -> list[str]:
            return sorted({str(row.get(field)).strip() for row in self.records if row.get(field) not in (None, "")})[:limit]

        fields: list[dict[str, Any]] = []
        if any(row.get("age_ma") is not None for row in self.records):
            fields.append({"field": "age_ma", "semantic_role": "event_time", "operators": ["between"], "value_type": "number"})
        if any(row.get("isotope_system") not in (None, "") for row in self.records):
            fields.append({"field": "isotope_system", "semantic_role": "clock_system", "operators": ["eq", "in"], "value_type": "string", "known_values": values("isotope_system")})
        for field, role in (("mineral", "measured_material"), ("rock_type", "material_class")):
            if any(row.get(field) not in (None, "") for row in self.records):
                fields.append({"field": field, "semantic_role": role, "operators": ["eq", "in"], "value_type": "string", "known_values": values(field)})
        if any(row.get(field) not in (None, "") for row in self.records for field in ("region", "province", "basin", "locality", "city")):
            fields.append({
                "field": "location_text", "semantic_role": "named_spatial_extent",
                "operators": ["contains_any"], "value_type": "string",
                "bound_columns": ["region", "province", "basin", "locality", "city"],
                "sample_values": sorted({str(row.get(field)).strip() for row in self.records for field in ("region", "province", "basin", "locality", "city") if row.get(field) not in (None, "")})[:80],
            })
        return {
            "query_contract": "observation_query/v1", "source_id": self.source_id,
            "version": self.version, "fields": fields,
            "failure_mode": "A scope with an unsupported field or operator is rejected; an unresolvable scope returns zero rows.",
        }

    def describe_discovery_coverage(self) -> dict[str, Any]:
        """Cheap catalogue-level coverage for Scout; never a question filter."""

        ages = [row["age_ma"] for row in self.records if row.get("age_ma") is not None]
        coordinates = [row["position"] for row in self.records if row.get("position", {}).get("latitude") is not None and row.get("position", {}).get("longitude") is not None]
        age_errors = [row["age_error_ma"] for row in self.records if row.get("age_error_ma") is not None]
        spatial_errors = [point["uncertainty_km"] for point in coordinates if point.get("uncertainty_km") is not None]
        named_fields = ("region", "province", "basin", "city", "locality")
        named_extent = {
            field: sorted({str(row[field]).strip() for row in self.records if row.get(field) not in (None, "")})[:80]
            for field in named_fields
        }
        return {
            "record_count": len(self.records),
            "time_coverage_ma": {"min": min(ages), "max": max(ages)} if ages else None,
            "named_spatial_coverage": {field: values for field, values in named_extent.items() if values},
            "coordinate_coverage": round(len(coordinates) / len(self.records), 4) if self.records else 0.0,
            "age_error_coverage": round(len(age_errors) / len(self.records), 4) if self.records else 0.0,
            "median_spatial_uncertainty_km": median(spatial_errors) if spatial_errors else None,
            "minimum_resolvable_migration_km": round(median(spatial_errors) * 2, 2) if spatial_errors else None,
            "coordinate_bbox": {
                "min_latitude": min(point["latitude"] for point in coordinates),
                "max_latitude": max(point["latitude"] for point in coordinates),
                "min_longitude": min(point["longitude"] for point in coordinates),
                "max_longitude": max(point["longitude"] for point in coordinates),
            } if coordinates else None,
        }

    def query(self, scope: dict[str, Any]) -> list[dict[str, Any]]:
        # A shared event/material/spatial registry may receive a legacy broad
        # scope.  Event time and isotope filters do not semantically apply to
        # independent position/material observations, whose rows often have no
        # clock value.  Canonical scopes already make this decision per source
        # in ScopeCompiler, so this compatibility path only preserves direct
        # adapter callers.
        effective_scope = scope
        if self.sensor_kind != "event_clock_sensor" and scope.get("_scope_contract") != "observation_scope/v1":
            effective_scope = {
                key: value for key, value in scope.items()
                if key not in {"min_age_ma", "max_age_ma", "age_range_ma", "isotope_system"}
            }
        return [dict(record) for record in self.records if _in_scope(record, effective_scope)]

    def load_workspace_handle(self, handle: dict[str, Any]) -> list[dict[str, Any]]:
        """Reopen a generic in-memory/JSON materialization by its typed handle.

        ``data.tool`` is allowed to use a generic fallback for small
        sources, but its result must still be consumable by numerical tools.
        The handle carries the already compiled source scope; older handles
        without that field are conservatively reconstructed from their typed
        predicates.  No fuzzy column matching is performed here.
        """

        if handle.get("source_id") != self.source_id:
            raise ValueError("Workspace handle belongs to a different source")
        scope = handle.get("scope")
        if not isinstance(scope, dict) or scope.get("_scope_contract") != "observation_scope/v1":
            scope = _scope_from_workspace_request(handle.get("request") or {}, self.source_id)
        rows = self.query(scope)
        expected = handle.get("record_count")
        if isinstance(expected, int) and expected >= 0 and len(rows) != expected:
            raise ValueError(f"Workspace handle record count changed: expected {expected}, loaded {len(rows)}")
        return rows

    def sample_records(self, scope: dict[str, Any], n: int = 5) -> list[dict[str, Any]]:
        return self.query(scope)[:n]

    def estimate_query(self, scope: dict[str, Any]) -> dict[str, Any]:
        count = len(self.query(scope))
        return {"estimated_records": count, "estimated_cost": "low" if count < 5_000 else "moderate", "query_scope": scope}

    def describe_resolution(self, scope: dict[str, Any]) -> dict[str, Any]:
        rows = self.query(scope)
        positions = [row["position"] for row in rows if row["position"]["latitude"] is not None]
        uncertainty = [p["uncertainty_km"] for p in positions if p.get("uncertainty_km") is not None]
        counts = Counter(p.get("resolution", "unresolved") for p in positions)
        coverage = len(positions) / len(rows) if rows else 0.0
        median_uncertainty = median(uncertainty) if uncertainty else None
        return {
            "coordinate_coverage": round(coverage, 4), "position_resolution_counts": dict(counts),
            "median_spatial_uncertainty_km": median_uncertainty,
            "minimum_resolvable_migration_km": round((median_uncertainty or 50.0) * 2, 2),
            "age_resolution": "event_age" if any(row.get("age_ma") is not None for row in rows) else "unavailable",
        }

    def describe_uncertainty(self, scope: dict[str, Any]) -> dict[str, Any]:
        rows = self.query(scope)
        age_errors = [row["age_error_ma"] for row in rows if row.get("age_error_ma") is not None]
        spatial = self.describe_resolution(scope).get("median_spatial_uncertainty_km")
        return {"age_error_coverage": round(len(age_errors) / len(rows), 4) if rows else 0.0, "median_age_error_ma": median(age_errors) if age_errors else None, "median_spatial_uncertainty_km": spatial}

    def profile_coverage(self, scope: dict[str, Any]) -> CoverageProfile:
        rows = self.query(scope)
        tracked = ("age_ma", "age_error_ma", "latitude", "longitude", "rock_type", "source_id", "event_interpretation", "mineral", "technique")
        missingness = {key: round(sum(row.get(key) in (None, "") for row in rows) / len(rows), 4) if rows else 1.0 for key in tracked}
        systems = Counter(row.get("isotope_system", "unknown") for row in rows)
        return CoverageProfile(
            self.source_id, len(rows), len(self.records), {"fields": sorted({key for row in rows for key in row}), "isotope_systems": dict(systems), "observables": self.describe_observables(scope)}, missingness,
            self.describe_resolution(scope), self.describe_uncertainty(scope), self.estimate_query(scope), self.sample_records(scope), self._warnings(rows), self.version,
        )

    def _warnings(self, rows: list[dict[str, Any]]) -> list[str]:
        if not rows:
            return ["当前 scope 没有匹配观测。"]
        warnings: list[str] = []
        spatial_coverage = sum(row["position"]["latitude"] is not None for row in rows) / len(rows)
        if spatial_coverage < .5:
            warnings.append("空间坐标覆盖不足；不能声明精细尺度迁移。")
        if sum(row.get("age_error_ma") is not None for row in rows) / len(rows) < .5:
            warnings.append("年龄误差覆盖不足；时间结构只能作保守解释。")
        return warnings


class JsonObservationSource(InMemoryObservationSource):
    def __init__(self, path: str | Path, *, source_id: str = "geochronology_json") -> None:
        file_path = Path(path)
        payload = json.loads(file_path.read_text(encoding="utf-8"))
        records = payload if isinstance(payload, list) else payload.get("records") or payload.get("data") or payload.get("rows") or []
        if not isinstance(records, list):
            raise ValueError("JSON source must contain a list or records/data/rows list")
        super().__init__(records, source_id=source_id, version=f"json:{file_path.name}:{file_path.stat().st_mtime_ns}")


class SpatialConstraintSource(InMemoryObservationSource):
    """Independent positions with provenance and σ, joined to Event Clock records by paper/sample."""

    def __init__(self, records: list[dict[str, Any]], *, source_id: str = "spatial_constraints", version: str = "memory") -> None:
        super().__init__(records, source_id=source_id, version=version, sensor_kind="spatial_constraint", name="Spatial constraint sensor", description="Independent position observations; records must include paper_id/source_id and sample_id for event joins.")


class MaterialStateSensorSource(InMemoryObservationSource):
    """Future Hf/Nd/geochemistry-style state observations joined by paper/sample identifiers."""

    def __init__(self, records: list[dict[str, Any]], *, source_id: str = "material_state", version: str = "memory") -> None:
        super().__init__(records, source_id=source_id, version=version, sensor_kind="material_state_sensor", name="Material-state sensor", description="Independent material-state observations; fields remain native and are attached to matching events.")


class OpportunisticSensorSource(InMemoryObservationSource):
    """A non-authoritative, optional acquisition result. It is never joined automatically."""

    def __init__(self, records: list[dict[str, Any]], *, source_id: str = "opportunistic", version: str = "transient", acquisition_status: str = "ok") -> None:
        super().__init__(records, source_id=source_id, version=version, sensor_kind="opportunistic_sensor", name="Opportunistic external sensor", description="Weak or temporary external evidence; must not be a main-chain dependency.")
        self.acquisition_status = acquisition_status

    def estimate_query(self, scope: dict[str, Any]) -> dict[str, Any]:
        result = super().estimate_query(scope)
        result["acquisition_status"] = self.acquisition_status
        return result


class TotalTreeXlsxSource(InMemoryObservationSource):
    """Semantic event adapter for the real reference/samples/observations/measure_events workbook."""

    def __init__(self, path: str | Path, *, source_id: str = "total_tree_geochronology") -> None:
        file_path = Path(path)
        try:
            import openpyxl  # type: ignore
        except ImportError as error:  # pragma: no cover - environment specific
            raise RuntimeError("TotalTreeXlsxSource requires the optional openpyxl dependency") from error
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        required = {"reference", "samples", "observations", "measure_events"}
        if not required.issubset(set(workbook.sheetnames)):
            raise ValueError(f"total_tree workbook must contain {sorted(required)}")
        tables = {name: self._table(workbook[name]) for name in required}
        samples = {(row.get("paper_id"), row.get("sample_id")): row for row in tables["samples"]}
        references = {row.get("paper_id"): row for row in tables["reference"]}
        spots: dict[tuple[Any, Any], list[dict[str, Any]]] = defaultdict(list)
        for row in tables["observations"]:
            spots[(row.get("paper_id"), row.get("sample_id"))].append(row)
        events: list[dict[str, Any]] = []
        for index, event in enumerate(tables["measure_events"], 1):
            key = (event.get("paper_id"), event.get("sample_id"))
            sample, reference, spot_rows = samples.get(key, {}), references.get(event.get("paper_id"), {}), spots.get(key, [])
            payload = {
                "event_payload": {key: event.get(key) for key in ("physical_method", "statistical_model", "distribution_state", "value_error_scale")},
                "observation_payload": {"spot_count": len(spot_rows), "fields": sorted({field for spot in spot_rows for field in spot}), "spots": spot_rows},
                "reference_payload": {key: reference.get(key) for key in ("title", "year", "doi")},
            }
            events.append({**sample, **event, **payload, "record_id": f"event:{index}", "source_id": event.get("paper_id")})
        manifest_id, bindings = load_semantic_manifest(Path(__file__).parent / "data_adapters" / "manifests" / "total_tree_event_clock.v1.json")
        super().__init__(events, source_id=source_id, version=f"xlsx:{file_path.name}:{file_path.stat().st_mtime_ns}", name="Total-tree geochronology event-clock sensor", description="Derived-event view over a read-only total_tree workbook; spot-level payload remains attached for QC tools.", semantic_bindings=bindings, semantic_manifest_id=manifest_id)

    @staticmethod
    def _table(sheet: Any) -> list[dict[str, Any]]:
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        return [{headers[index]: value for index, value in enumerate(row) if headers[index]} for row in rows]
